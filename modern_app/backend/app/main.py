from __future__ import annotations

from datetime import datetime
from pathlib import Path
from tempfile import gettempdir, mkdtemp
from typing import Any, Literal
from contextlib import closing
import csv
import hmac
import io
import json
import logging
import os
import socket
import sys
import time
from uuid import uuid4

import sqlite3
from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from starlette.background import BackgroundTask

from finance_app.exporter import export_date_range_report, export_filtered_movimientos, export_monthly_report, export_yearly_report
from finance_app.db import CURRENT_SCHEMA_VERSION
from finance_app.services import (
    LOCAL_OWNER_ID,
    FinanceService,
    GastoFijoInput,
    GastoProgramadoInput,
    MetaAhorroInput,
    MovimientoInput,
    PresupuestoInput,
    TagInput,
    get_current_owner_id,
    normalize_owner_id,
    reset_current_owner_id,
    set_current_owner_id,
)
from finance_app.paths import get_app_data_dir, get_backup_dir, get_data_dir, get_db_path, get_logs_dir
from openpyxl import load_workbook

from .deps import ensure_app_data_initialized, get_database_readiness, get_last_init_status, get_service, invalidate_app_data_initialized
from .settings import ORIGINAL_DB_PATH, WEB_DB_PATH
from .schemas import BackupFrequencyIn, BackupRestoreIn, BackupRestorePathIn, CategoriaIn, GastoFijoIn, GastoProgramadoIn, MetaAhorroIn, MovimientoIn, PresupuestoIn, TagIn

_IS_FROZEN = bool(getattr(sys, "frozen", False))
# La documentacion interactiva ayuda en desarrollo, pero en la app instalada
# expondria superficie interna innecesaria aunque los endpoints tengan token.
app = FastAPI(
    title="Registro Finanzas API",
    version="3.0.1",
    docs_url=None if _IS_FROZEN else "/docs",
    redoc_url=None if _IS_FROZEN else "/redoc",
    openapi_url=None if _IS_FROZEN else "/openapi.json",
)
_LOCAL_TOKEN_HEADER = "X-Scisonomics-Local-Token"
_LOCAL_TOKEN = os.getenv("SCISONOMICS_LOCAL_TOKEN", "").strip()
_DEV_MODE_WITHOUT_LOCAL_TOKEN = not _LOCAL_TOKEN and not _IS_FROZEN
_PUBLIC_PATHS = {"/health", "/ready"} | (set() if _IS_FROZEN else {"/openapi.json"})
_PUBLIC_PREFIXES = () if _IS_FROZEN else ("/docs", "/redoc")

_LOG_FILE = get_logs_dir() / "backend-startup.log"
_logger = logging.getLogger("scisonomics.backend")
if not _logger.handlers:
    _logger.setLevel(logging.INFO)
    _handler: logging.Handler | None = None
    for candidate in (
        _LOG_FILE,
        get_app_data_dir() / "logs" / "backend-startup.log",
        Path.cwd() / "backend-startup.log",
    ):
        try:
            candidate.parent.mkdir(parents=True, exist_ok=True)
            _handler = logging.FileHandler(candidate, encoding="utf-8")
            break
        except OSError:
            continue
    if _handler is None:
        _handler = logging.StreamHandler()
    _handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    _logger.addHandler(_handler)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://127.0.0.1:3000",
        "http://127.0.0.1",
        "http://localhost:3000",
        "http://localhost",
        "http://tauri.localhost",
        "https://tauri.localhost",
        "tauri://localhost",
        "asset://localhost",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type", "X-Scisonomics-Owner-Id", _LOCAL_TOKEN_HEADER, "Authorization"],
)


@app.middleware("http")
async def owner_context_middleware(request: Request, call_next):
    if request.method != "OPTIONS" and not _is_public_path(request.url.path):
        try:
            _require_local_token(request)
        except HTTPException as exc:
            return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})
    token = set_current_owner_id(request.headers.get("X-Scisonomics-Owner-Id"))
    try:
        return await call_next(request)
    finally:
        reset_current_owner_id(token)


def _is_public_path(path: str) -> bool:
    return path in _PUBLIC_PATHS or any(path.startswith(prefix) for prefix in _PUBLIC_PREFIXES)


def _require_local_token(request: Request) -> None:
    if _DEV_MODE_WITHOUT_LOCAL_TOKEN:
        return
    if not _LOCAL_TOKEN:
        raise HTTPException(status_code=503, detail="El servicio local no tiene token de seguridad configurado.")
    provided = request.headers.get(_LOCAL_TOKEN_HEADER, "")
    if not provided or not hmac.compare_digest(provided, _LOCAL_TOKEN):
        raise HTTPException(status_code=401, detail="No se pudo validar la conexión local de ScisoNomics.")


def _require_debug_endpoints_enabled() -> None:
    enabled = os.getenv("SCISONOMICS_ENABLE_DEBUG_ENDPOINTS", "").strip().lower() in {"1", "true", "yes", "on"}
    if not enabled:
        raise HTTPException(status_code=404, detail="Not Found")


def _current_owner() -> str:
    return get_current_owner_id()


def _safe_owner_for_log(owner: str) -> str:
    # El owner completo no es necesario para soporte y puede identificar una cuenta.
    return LOCAL_OWNER_ID if owner == LOCAL_OWNER_ID else f"{owner[:6]}..."


def _service_value_error(exc: ValueError, resource_label: str = "El recurso") -> HTTPException:
    # Un ID ajeno y un ID inexistente deben verse igual desde la API: responder
    # 404 evita filtrar existencia de datos pertenecientes a otra cuenta.
    if "no existe o no pertenece a la cuenta activa" in str(exc).lower():
        return HTTPException(status_code=404, detail=f"{resource_label} no se encontro o no pertenece a la cuenta activa.")
    return HTTPException(status_code=400, detail=str(exc))


def _cleanup_temp_snapshot(path: Path) -> None:
    # Los exports usan carpetas temporales dedicadas. Limpiarlas luego de enviar
    # el archivo evita acumular snapshots sin tocar backups persistentes.
    try:
        path.unlink(missing_ok=True)
        path.parent.rmdir()
    except OSError:
        _logger.warning("No se pudo limpiar snapshot temporal. file=%s", path.name)


def _cleanup_stale_temp_snapshots(max_age_seconds: int = 600) -> None:
    # BackgroundTask no corre si el streaming se aborta. Recolectar solo
    # carpetas temporales propias y antiguas cubre ese caso sin tocar backups.
    temp_root = Path(gettempdir()).resolve()
    cutoff = time.time() - max_age_seconds
    for prefix in ("finanzas_backup_", "finanzas_download_"):
        for candidate in temp_root.glob(f"{prefix}*"):
            try:
                folder = candidate.resolve()
                if folder.parent != temp_root or not folder.is_dir() or folder.stat().st_mtime > cutoff:
                    continue
                for child in folder.iterdir():
                    if child.is_file():
                        child.unlink(missing_ok=True)
                folder.rmdir()
            except OSError:
                _logger.warning("No se pudo limpiar snapshot temporal antiguo. folder=%s", candidate.name)


def _require_cloud_owner() -> str:
    owner = _current_owner()
    if not owner or owner == LOCAL_OWNER_ID:
        raise HTTPException(status_code=400, detail="No hay una cuenta activa para sincronizar.")
    return owner


def _owner_filter(alias: str | None = None) -> str:
    prefix = f"{alias}." if alias else ""
    return f"{prefix}owner_user_id = ?"


@app.exception_handler(FileNotFoundError)
async def file_not_found_handler(_: Request, exc: FileNotFoundError):
    return JSONResponse(status_code=500, content={"detail": str(exc) or "Error de archivo local."})


@app.on_event("startup")
def log_db_path() -> None:
    try:
        _cleanup_stale_temp_snapshots()
        ensure_app_data_initialized()
        exists = WEB_DB_PATH.exists()
        size_bytes = WEB_DB_PATH.stat().st_size if exists else 0
        _logger.info("Startup backend OK. db_file=%s exists=%s size=%s", WEB_DB_PATH.name, exists, size_bytes)
    except Exception as exc:
        _logger.exception("Error inicializando backend local. error_type=%s", type(exc).__name__)


@app.get("/health")
def health():
    try:
        ensure_app_data_initialized()
        return {"ok": True, "status": "healthy", "version": app.version}
    except Exception as exc:
        _logger.exception("Error en /health. error_type=%s", type(exc).__name__)
        return JSONResponse(
            status_code=500,
            content={"ok": False, "status": "unavailable", "version": app.version},
        )


@app.get("/ready")
def ready():
    status = get_database_readiness()
    if status.get("database_ready"):
        return {
            "ok": True,
            "status": "ready",
            "database_ready": True,
            "initializing": False,
            "version": app.version,
        }
    return JSONResponse(
        status_code=503,
        content={
            "ok": False,
            "status": "initializing" if status.get("initializing", False) else "unavailable",
            "database_ready": False,
            "initializing": bool(status.get("initializing", False)),
            "version": app.version,
        },
    )


def _safe_app_paths() -> dict[str, str]:
    data_dir = get_data_dir()
    backups_dir = get_backup_dir()
    logs_dir = get_logs_dir()
    for folder in (data_dir, backups_dir, logs_dir):
        try:
            folder.mkdir(parents=True, exist_ok=True)
        except OSError:
            _logger.warning("No se pudo asegurar carpeta de app. folder=%s", folder.name)
    return {
        "database_path": str(WEB_DB_PATH),
        "data_dir": str(data_dir),
        "backups_path": str(backups_dir),
        "logs_path": str(logs_dir),
    }


@app.get("/app/paths")
def app_paths():
    ensure_app_data_initialized()
    return {"ok": True, "version": app.version, **_safe_app_paths()}


@app.get("/app/diagnostics")
def app_diagnostics():
    try:
        ensure_app_data_initialized()
    except Exception:
        _logger.exception("Error preparando diagnostico local.")
    status = get_last_init_status()
    paths = _safe_app_paths()
    return {
        "ok": bool(status.get("database_ready", False)),
        "version": app.version,
        "database_ready": bool(status.get("database_ready", False)),
        "initializing": bool(status.get("initializing", False)),
        "database_error": status.get("database_error"),
        "db_exists": bool(Path(paths["database_path"]).exists()),
        "frozen": bool(getattr(sys, "frozen", False)),
        **paths,
    }


@app.get("/local/auth-check")
def local_auth_check():
    return {"ok": True, "service": "local", "owner_user_id": _current_owner()}


@app.get("/meta")
def meta(service: FinanceService = Depends(get_service)):
    return {
        "database": {
            "working_copy": str(service.db.db_path),
            "original": str(ORIGINAL_DB_PATH),
        }
    }


@app.get("/debug/db-path", dependencies=[Depends(_require_debug_endpoints_enabled)])
def debug_db_path(service: FinanceService = Depends(get_service)):
    _require_debug_endpoints_enabled()
    db_path = Path(service.db.db_path).resolve()
    exists = db_path.exists()
    size_bytes = db_path.stat().st_size if exists else 0
    counts = {
        "movimientos_count": 0,
        "categorias_count": 0,
        "gastos_fijos_count": 0,
        "gastos_programados_count": 0,
    }
    if exists:
        with service.db.connect() as conn:
            counts["movimientos_count"] = int(conn.execute("SELECT COUNT(*) FROM movimientos").fetchone()[0])
            counts["categorias_count"] = int(conn.execute("SELECT COUNT(*) FROM categorias").fetchone()[0])
            counts["gastos_fijos_count"] = int(conn.execute("SELECT COUNT(*) FROM gastos_fijos").fetchone()[0])
            counts["gastos_programados_count"] = int(conn.execute("SELECT COUNT(*) FROM gastos_programados").fetchone()[0])
    return {
        "db_path": str(db_path),
        "exists": exists,
        "size_bytes": size_bytes,
        **counts,
    }


@app.get("/movimientos")
def list_movimientos(
    month: int | None = Query(default=None, ge=1, le=12),
    year: int | None = Query(default=None),
    tipo: Literal["todos", "ingreso", "gasto", "inversion", "ahorro"] = "todos",
    search: str = "",
    categoria: str = "",
    min_monto: float | None = Query(default=None, ge=0),
    max_monto: float | None = Query(default=None, ge=0),
    service: FinanceService = Depends(get_service),
):
    now = datetime.now()
    month = month if month is not None else now.month
    year = year if year is not None else now.year
    current_type = None if tipo in ("todos", "inversion") else tipo
    rows = service.list_movimientos(month, year, current_type)
    if tipo == "inversion":
        rows = [r for r in rows if "invers" in str(r.get("categoria", "")).lower()]
    term = search.strip().lower()
    category_term = categoria.strip().lower()
    if term:
        rows = [
            row
            for row in rows
            if term in str(row.get("descripcion", "")).lower()
            or term in str(row.get("categoria", "")).lower()
            or term in str(row.get("fecha", "")).lower()
        ]
    if category_term:
        rows = [row for row in rows if category_term in str(row.get("categoria", "")).lower()]
    if min_monto is not None:
        rows = [row for row in rows if float(row.get("monto", 0) or 0) >= min_monto]
    if max_monto is not None:
        rows = [row for row in rows if float(row.get("monto", 0) or 0) <= max_monto]

    summary = service.get_resumen_mensual_con_saldo(month, year)
    visible_total = sum(float(r["monto"]) if r["tipo"] == "ingreso" else -float(r["monto"]) for r in rows)
    return {
        "rows": rows,
        "summary": summary,
        "visible_count": len(rows),
        "visible_total": visible_total,
    }


@app.post("/movimientos")
def create_movimiento(payload: MovimientoIn, service: FinanceService = Depends(get_service)):
    try:
        service.create_movimiento(MovimientoInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El movimiento") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.put("/movimientos/{movimiento_id}")
def update_movimiento(movimiento_id: int, payload: MovimientoIn, service: FinanceService = Depends(get_service)):
    try:
        service.update_movimiento(movimiento_id, MovimientoInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El movimiento") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/movimientos/{movimiento_id}")
def delete_movimiento(movimiento_id: int, service: FinanceService = Depends(get_service)):
    try:
        service.delete_movimiento(movimiento_id)
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El movimiento") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/categorias")
def list_categorias(tipo: Literal["ingreso", "gasto", "ahorro", "inversion", "todos"] = "todos", service: FinanceService = Depends(get_service)):
    return service.list_categorias(None if tipo == "todos" else tipo)


@app.post("/categorias")
def create_categoria(payload: CategoriaIn, service: FinanceService = Depends(get_service)):
    try:
        service.create_categoria(payload.nombre, payload.tipo)
        return {"ok": True}
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.put("/categorias/{categoria_id}")
def update_categoria(categoria_id: int, payload: CategoriaIn, service: FinanceService = Depends(get_service)):
    try:
        service.update_categoria(categoria_id, payload.nombre, payload.tipo)
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "La categoria") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/categorias/{categoria_id}")
def delete_categoria(categoria_id: int, service: FinanceService = Depends(get_service)):
    try:
        service.delete_categoria(categoria_id)
        return {"ok": True}
    except ValueError as exc:
        message = str(exc)
        if "no existe o no pertenece a la cuenta activa" in message.lower():
            raise _service_value_error(exc, "La categoria") from exc
        if "movimientos asociados" in message.lower():
            raise HTTPException(status_code=409, detail="No se puede eliminar esta categoría porque tiene movimientos asociados.") from exc
        if "no se encontro la categoria" in message.lower() or "no se encontró la categoria" in message.lower():
            raise HTTPException(status_code=404, detail="No se encontró la categoría.") from exc
        raise HTTPException(status_code=400, detail=message) from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/gastos-fijos")
def list_gastos_fijos(service: FinanceService = Depends(get_service)):
    return service.list_gastos_fijos()


@app.post("/gastos-fijos")
def create_gasto_fijo(payload: GastoFijoIn, service: FinanceService = Depends(get_service)):
    try:
        service.create_gasto_fijo(GastoFijoInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El gasto fijo") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.put("/gastos-fijos/{gasto_id}")
def update_gasto_fijo(gasto_id: int, payload: GastoFijoIn, service: FinanceService = Depends(get_service)):
    try:
        service.update_gasto_fijo(gasto_id, GastoFijoInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El gasto fijo") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/gastos-fijos/{gasto_id}")
def delete_gasto_fijo(gasto_id: int, service: FinanceService = Depends(get_service)):
    try:
        service.delete_gasto_fijo(gasto_id)
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El gasto fijo") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/gastos-programados")
def list_gastos_programados(
    estado: Literal["pendiente", "pagado", "cancelado", "todos"] = "todos",
    dias: int | None = Query(default=None, ge=1),
    service: FinanceService = Depends(get_service),
):
    filtro_estado = None if estado == "todos" else estado
    rows = service.get_gastos_programados(filtro_estado, dias)
    return rows


@app.post("/gastos-programados")
def create_gasto_programado(payload: GastoProgramadoIn, service: FinanceService = Depends(get_service)):
    try:
        service.create_gasto_programado(GastoProgramadoInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El gasto programado") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.put("/gastos-programados/{gasto_id}")
def update_gasto_programado(gasto_id: int, payload: GastoProgramadoIn, service: FinanceService = Depends(get_service)):
    try:
        service.update_gasto_programado(gasto_id, GastoProgramadoInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El gasto programado") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/gastos-programados/{gasto_id}")
def delete_gasto_programado(gasto_id: int, service: FinanceService = Depends(get_service)):
    try:
        service.delete_gasto_programado(gasto_id)
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El gasto programado") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/gastos-programados/{gasto_id}/marcar-pagado")
def marcar_pagado(gasto_id: int, service: FinanceService = Depends(get_service)):
    try:
        return service.marcar_gasto_programado_pagado(gasto_id)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/estadisticas")
def get_stats(month: int = Query(..., ge=1, le=12), year: int = Query(...), service: FinanceService = Depends(get_service)):
    summary = service.get_resumen_mensual_con_saldo(month, year)
    expenses_by_category = service.get_expenses_by_category(month, year, "gasto")
    trend = service.get_monthly_trend(year)
    plan = service.get_planificacion_resumen(month, year)
    month_totals = service.get_month_summary(month, year, None)
    return {
        "summary": summary,
        "month_totals": month_totals,
        "expenses_by_category": expenses_by_category,
        "trend": trend,
        "planificacion": plan,
    }


@app.get("/estadisticas/anual")
def get_stats_anual(year: int = Query(...), service: FinanceService = Depends(get_service)):
    with service.db.connect() as conn:
        totals_row = conn.execute(
            """
            SELECT
              COALESCE(SUM(CASE WHEN tipo='ingreso' THEN monto ELSE 0 END), 0) AS ingresos,
              COALESCE(SUM(CASE WHEN tipo='gasto' THEN monto ELSE 0 END), 0) AS gastos,
              COALESCE(SUM(CASE WHEN tipo='ahorro' THEN monto ELSE 0 END), 0) AS ahorros,
              COALESCE(SUM(CASE WHEN tipo='inversion' THEN monto ELSE 0 END), 0) AS inversiones,
              COUNT(*) AS movimientos
            FROM movimientos
            WHERE strftime('%Y', fecha) = ?
              AND (deleted_at IS NULL OR deleted_at = '')
              AND owner_user_id = ?
            """,
            (str(year), _current_owner()),
        ).fetchone()

        monthly_rows = conn.execute(
            """
            SELECT
              CAST(strftime('%m', fecha) AS INTEGER) AS mes,
              COALESCE(SUM(CASE WHEN tipo='ingreso' THEN monto ELSE 0 END), 0) AS ingresos,
              COALESCE(SUM(CASE WHEN tipo='gasto' THEN monto ELSE 0 END), 0) AS gastos,
              COALESCE(SUM(CASE WHEN tipo='ahorro' THEN monto ELSE 0 END), 0) AS ahorros,
              COALESCE(SUM(CASE WHEN tipo='inversion' THEN monto ELSE 0 END), 0) AS inversiones
            FROM movimientos
            WHERE strftime('%Y', fecha) = ?
              AND (deleted_at IS NULL OR deleted_at = '')
              AND owner_user_id = ?
            GROUP BY strftime('%m', fecha)
            ORDER BY mes
            """,
            (str(year), _current_owner()),
        ).fetchall()

        categories_rows = conn.execute(
            """
            SELECT c.nombre AS categoria, COALESCE(SUM(m.monto),0) AS total, COUNT(m.id) AS movimientos
            FROM movimientos m
            JOIN categorias c ON c.id = m.categoria_id AND c.owner_user_id = m.owner_user_id
            WHERE strftime('%Y', m.fecha) = ? AND m.tipo = 'gasto'
              AND (m.deleted_at IS NULL OR m.deleted_at = '')
              AND (c.deleted_at IS NULL OR c.deleted_at = '')
              AND m.owner_user_id = ?
              AND c.owner_user_id = ?
            GROUP BY c.nombre
            HAVING total > 0
            ORDER BY total DESC
            """,
            (str(year), _current_owner(), _current_owner()),
        ).fetchall()

    by_month = {
        int(row["mes"]): {
            "mes": int(row["mes"]),
            "ingresos": float(row["ingresos"] or 0),
            "gastos": float(row["gastos"] or 0),
            "ahorros": float(row["ahorros"] or 0),
            "inversiones": float(row["inversiones"] or 0),
        }
        for row in monthly_rows
    }
    monthly = []
    for m in range(1, 13):
        item = by_month.get(m, {"mes": m, "ingresos": 0.0, "gastos": 0.0, "ahorros": 0.0, "inversiones": 0.0})
        item["balance"] = item["ingresos"] - item["gastos"] - item["ahorros"] - item["inversiones"]
        monthly.append(item)

    total_ingresos = float(totals_row["ingresos"] or 0)
    total_gastos = float(totals_row["gastos"] or 0)
    total_ahorros = float(totals_row["ahorros"] or 0)
    total_inversiones = float(totals_row["inversiones"] or 0)
    balance_anual = total_ingresos - total_gastos - total_ahorros - total_inversiones
    movimientos_total = int(totals_row["movimientos"] or 0)

    mes_mayor_gasto = max(monthly, key=lambda x: x["gastos"]) if monthly else None
    mes_mayor_ingreso = max(monthly, key=lambda x: x["ingresos"]) if monthly else None
    categoria_mayor_gasto = dict(categories_rows[0]) if categories_rows else None

    return {
        "year": year,
        "totals": {
            "ingresos": total_ingresos,
            "gastos": total_gastos,
            "ahorros": total_ahorros,
            "inversiones": total_inversiones,
            "balance": balance_anual,
            "movimientos": movimientos_total,
        },
        "promedios_mensuales": {
            "ingresos": total_ingresos / 12.0,
            "gastos": total_gastos / 12.0,
            "balance": balance_anual / 12.0,
        },
        "mes_mayor_gasto": mes_mayor_gasto,
        "mes_mayor_ingreso": mes_mayor_ingreso,
        "categoria_mayor_gasto": categoria_mayor_gasto,
        "monthly": monthly,
        "gastos_por_categoria": [dict(r) for r in categories_rows],
    }


@app.get("/resumen-mensual")
def resumen_mensual(
    month: int | None = Query(default=None, ge=1, le=12),
    year: int | None = Query(default=None),
    mes: int | None = Query(default=None, ge=1, le=12),
    anio: int | None = Query(default=None),
    service: FinanceService = Depends(get_service),
):
    month = month if month is not None else mes
    year = year if year is not None else anio
    if month is None or year is None:
        raise HTTPException(status_code=400, detail="Debes indicar month/year o mes/anio.")
    return service.get_resumen_mensual_potente(month, year)


@app.get("/presupuestos")
def list_presupuestos(month: int = Query(..., ge=1, le=12), year: int = Query(...), service: FinanceService = Depends(get_service)):
    return service.list_presupuestos(month, year)


@app.post("/presupuestos")
def upsert_presupuesto(payload: PresupuestoIn, service: FinanceService = Depends(get_service)):
    try:
        service.create_or_update_presupuesto(PresupuestoInput(**payload.model_dump()))
        return {"ok": True}
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/presupuestos/{presupuesto_id}")
def delete_presupuesto(presupuesto_id: int, service: FinanceService = Depends(get_service)):
    try:
        service.delete_presupuesto(presupuesto_id)
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "El presupuesto") from exc
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/settings/info")
def settings_info(service: FinanceService = Depends(get_service)):
    ensure_app_data_initialized()
    status = get_last_init_status()
    db_path = Path(str(status.get("db_path", service.db.db_path)))
    data_dir = Path(str(status.get("data_dir", get_data_dir())))
    backups_dir = Path(str(status.get("backups_dir", get_backup_dir())))
    logs_dir = Path(str(status.get("logs_dir", get_logs_dir())))
    exists = db_path.exists()
    size_bytes = db_path.stat().st_size if exists else 0
    counts = {"movimientos": 0, "categorias": 0, "presupuestos": 0, "metas": 0}
    if exists:
        with service.db.connect() as conn:
            counts["movimientos"] = int(conn.execute("SELECT COUNT(*) FROM movimientos WHERE owner_user_id = ?", (_current_owner(),)).fetchone()[0])
            counts["categorias"] = int(conn.execute("SELECT COUNT(*) FROM categorias WHERE owner_user_id = ?", (_current_owner(),)).fetchone()[0])
            counts["presupuestos"] = int(conn.execute("SELECT COUNT(*) FROM presupuestos WHERE owner_user_id = ?", (_current_owner(),)).fetchone()[0])
            counts["metas"] = int(conn.execute("SELECT COUNT(*) FROM metas_ahorro WHERE owner_user_id = ?", (_current_owner(),)).fetchone()[0])
    return {
        "version": app.version,
        "backend_ok": True,
        "db_path": str(db_path),
        "db_exists": exists,
        "db_size": size_bytes,
        "data_dir": str(data_dir),
        "backups_dir": str(backups_dir),
        "logs_dir": str(logs_dir),
        "logs_exists": logs_dir.exists(),
        "app_data_dir": str(get_app_data_dir()),
        "migrations_status": "ok",
        "db_initialized": bool(status.get("db_initialized", False)),
        "frozen": bool(getattr(sys, "frozen", False)),
        "executable": str(sys.executable),
        "counts": counts,
    }


@app.get("/sync/status")
def sync_status(service: FinanceService = Depends(get_service)):
    ensure_app_data_initialized()
    owner = _current_owner()
    candidate_tables = [
        "movimientos",
        "categorias",
        "presupuestos",
        "metas_ahorro",
        "gastos_fijos",
        "gastos_programados",
        "tags",
        "movimiento_tags",
    ]
    tables: dict[str, dict[str, int | bool]] = {}
    with service.db.connect() as conn:
        existing_tables = {
            str(row["name"])
            for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()
        }
        for table in candidate_tables:
            if table not in existing_tables:
                continue
            columns = {
                str(row["name"])
                for row in conn.execute(f"PRAGMA table_info({table})").fetchall()
            }
            has_sync_columns = {"sync_id", "sync_status", "deleted_at"}.issubset(columns)
            if not has_sync_columns:
                tables[table] = {
                    "total": int(conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]),
                    "pending": 0,
                    "synced": 0,
                    "deleted": 0,
                    "missing_sync_id": 0,
                    "sync_ready": False,
                }
                continue
            row = conn.execute(
                f"""
                SELECT
                  COUNT(*) AS total,
                  COALESCE(SUM(CASE WHEN sync_status = 'pending' THEN 1 ELSE 0 END), 0) AS pending,
                  COALESCE(SUM(CASE WHEN sync_status = 'synced' THEN 1 ELSE 0 END), 0) AS synced,
                  COALESCE(SUM(CASE WHEN sync_status = 'deleted' OR deleted_at IS NOT NULL THEN 1 ELSE 0 END), 0) AS deleted,
                  COALESCE(SUM(CASE WHEN sync_id IS NULL OR trim(sync_id) = '' THEN 1 ELSE 0 END), 0) AS missing_sync_id
                FROM {table}
                WHERE owner_user_id = ?
                """
                ,
                (owner,),
            ).fetchone()
            tables[table] = {
                "total": int(row["total"] or 0),
                "pending": int(row["pending"] or 0),
                "synced": int(row["synced"] or 0),
                "deleted": int(row["deleted"] or 0),
                "missing_sync_id": int(row["missing_sync_id"] or 0),
                "sync_ready": int(row["missing_sync_id"] or 0) == 0,
            }

    return {
        "ok": True,
        "sync_ready": all(bool(info.get("sync_ready")) for info in tables.values()),
        "tables": tables,
    }


def _ensure_sync_history_table(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS sync_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sync_id TEXT NOT NULL UNIQUE,
            device_id TEXT,
            owner_user_id TEXT DEFAULT 'local',
            mode TEXT NOT NULL CHECK(mode IN ('manual', 'auto')),
            status TEXT NOT NULL CHECK(status IN ('success', 'error', 'skipped')),
            started_at TEXT NOT NULL,
            finished_at TEXT,
            duration_ms INTEGER DEFAULT 0,
            pending_total INTEGER DEFAULT 0,
            pushed_total INTEGER DEFAULT 0,
            pulled_total INTEGER DEFAULT 0,
            deleted_total INTEGER DEFAULT 0,
            conflicts_total INTEGER DEFAULT 0,
            remote_changes_total INTEGER DEFAULT 0,
            applied_remote_total INTEGER DEFAULT 0,
            kept_local_total INTEGER DEFAULT 0,
            error_message TEXT,
            details_json TEXT
        );
        CREATE TABLE IF NOT EXISTS sync_conflicts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conflict_id TEXT NOT NULL UNIQUE,
            owner_user_id TEXT DEFAULT 'local',
            table_name TEXT NOT NULL,
            record_sync_id TEXT NOT NULL,
            local_updated_at TEXT,
            remote_updated_at TEXT,
            last_synced_at TEXT,
            resolution TEXT NOT NULL CHECK(resolution IN ('kept_local', 'applied_remote', 'ignored')),
            remote_device_id TEXT,
            remote_device_name TEXT,
            detected_at TEXT NOT NULL,
            resolved_at TEXT,
            details_json TEXT
        );
        CREATE INDEX IF NOT EXISTS idx_sync_history_finished_at ON sync_history(finished_at);
        CREATE INDEX IF NOT EXISTS idx_sync_history_status ON sync_history(status);
        CREATE INDEX IF NOT EXISTS idx_sync_conflicts_detected_at ON sync_conflicts(detected_at);
        CREATE INDEX IF NOT EXISTS idx_sync_conflicts_record ON sync_conflicts(table_name, record_sync_id);
        """
    )
    for column in ("conflicts_total", "remote_changes_total", "applied_remote_total", "kept_local_total"):
        _ensure_column(conn, "sync_history", column, "INTEGER DEFAULT 0")
    _ensure_column(conn, "sync_history", "owner_user_id", "TEXT DEFAULT 'local'")
    _ensure_column(conn, "sync_conflicts", "owner_user_id", "TEXT DEFAULT 'local'")
    conn.execute("UPDATE sync_history SET owner_user_id = 'local' WHERE owner_user_id IS NULL OR trim(owner_user_id) = ''")
    conn.execute("UPDATE sync_conflicts SET owner_user_id = 'local' WHERE owner_user_id IS NULL OR trim(owner_user_id) = ''")


def _ensure_column(conn: sqlite3.Connection, table: str, column: str, definition: str) -> None:
    columns = {str(row["name"]) for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}
    if column not in columns:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def _config_get(conn: sqlite3.Connection, key: str) -> str | None:
    row = conn.execute("SELECT value FROM app_config WHERE key = ?", (key,)).fetchone()
    return str(row["value"]) if row and row["value"] is not None else None


def _config_set(conn: sqlite3.Connection, key: str, value: str) -> None:
    conn.execute(
        """
        INSERT INTO app_config (key, value, updated_at)
        VALUES (?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(key) DO UPDATE SET value = excluded.value, updated_at = CURRENT_TIMESTAMP
        """,
        (key, value),
    )


def _get_or_create_device_info(conn: sqlite3.Connection) -> dict[str, str]:
    device_id = _config_get(conn, "device_id")
    if not device_id:
        device_id = str(uuid4())
        _config_set(conn, "device_id", device_id)

    device_name = _config_get(conn, "device_name")
    if not device_name:
        device_name = socket.gethostname().strip() or "Este dispositivo"
        _config_set(conn, "device_name", device_name)

    return {"device_id": device_id, "device_name": device_name}


def _sync_status_tables(conn: sqlite3.Connection, owner: str | None = None) -> dict[str, dict[str, int]]:
    owner = normalize_owner_id(owner or _current_owner())
    tables: dict[str, dict[str, int]] = {}
    for table in SYNC_TABLES:
        row = conn.execute(
            f"""
            SELECT
              COUNT(*) AS total,
              COALESCE(SUM(CASE WHEN sync_status = 'pending' THEN 1 ELSE 0 END), 0) AS pending,
              COALESCE(SUM(CASE WHEN sync_status = 'pending' AND deleted_at IS NOT NULL AND deleted_at <> '' THEN 1 ELSE 0 END), 0) AS deleted_pending,
              COALESCE(SUM(CASE WHEN sync_id IS NULL OR trim(sync_id) = '' THEN 1 ELSE 0 END), 0) AS missing_sync_id
            FROM {table}
            WHERE owner_user_id = ?
            """
            ,
            (owner,),
        ).fetchone()
        tables[table] = {
            "total": int(row["total"] or 0),
            "pending": int(row["pending"] or 0),
            "deleted_pending": int(row["deleted_pending"] or 0),
            "missing_sync_id": int(row["missing_sync_id"] or 0),
        }
    return tables


def _last_remote_change_at(conn: sqlite3.Connection, owner: str | None = None) -> str | None:
    owner = normalize_owner_id(owner or _current_owner())
    values: list[str] = []
    for table in SYNC_TABLES:
        row = conn.execute(
            f"SELECT MAX(last_remote_updated_at) AS value FROM {table} WHERE last_remote_updated_at IS NOT NULL AND last_remote_updated_at <> '' AND owner_user_id = ?",
            (owner,),
        ).fetchone()
        if row and row["value"]:
            values.append(str(row["value"]))
    return max(values) if values else None


def _ensure_sync_metadata_columns(conn: sqlite3.Connection) -> None:
    for table in SYNC_TABLES:
        _ensure_column(conn, table, "owner_user_id", "TEXT DEFAULT 'local'")
        conn.execute(f"UPDATE {table} SET owner_user_id = 'local' WHERE owner_user_id IS NULL OR trim(owner_user_id) = ''")
        _ensure_column(conn, table, "last_remote_device_id", "TEXT")
        _ensure_column(conn, table, "last_remote_device_name", "TEXT")
        _ensure_column(conn, table, "last_remote_updated_at", "TEXT")


def _history_row(row: sqlite3.Row | None) -> dict[str, Any] | None:
    if row is None:
        return None
    item = dict(row)
    raw_details = item.pop("details_json", None)
    if raw_details:
        try:
            item["details"] = json.loads(raw_details)
        except json.JSONDecodeError:
            item["details"] = None
    else:
        item["details"] = None
    return item


def _latest_history(conn: sqlite3.Connection, status: str, owner: str | None = None) -> dict[str, Any] | None:
    _ensure_sync_history_table(conn)
    owner = normalize_owner_id(owner or _current_owner())
    row = conn.execute(
        """
        SELECT sync_id, device_id, mode, status, started_at, finished_at, duration_ms,
               pending_total, pushed_total, pulled_total, deleted_total,
               conflicts_total, remote_changes_total, applied_remote_total, kept_local_total,
               error_message, details_json
        FROM sync_history
        WHERE status = ? AND owner_user_id = ?
        ORDER BY COALESCE(finished_at, started_at) DESC, id DESC
        LIMIT 1
        """,
        (status, owner),
    ).fetchone()
    return _history_row(row)


def _conflict_row(row: sqlite3.Row | None) -> dict[str, Any] | None:
    if row is None:
        return None
    item = dict(row)
    raw_details = item.pop("details_json", None)
    if raw_details:
        try:
            item["details"] = json.loads(raw_details)
        except json.JSONDecodeError:
            item["details"] = None
    else:
        item["details"] = None
    return item


def _conflict_summary(conn: sqlite3.Connection, owner: str | None = None) -> dict[str, Any]:
    _ensure_sync_history_table(conn)
    owner = normalize_owner_id(owner or _current_owner())
    total = int(conn.execute("SELECT COUNT(*) FROM sync_conflicts WHERE owner_user_id = ?", (owner,)).fetchone()[0] or 0)
    recent = int(
        conn.execute(
            """
            SELECT COUNT(*) FROM sync_conflicts
            WHERE datetime(detected_at) >= datetime('now', '-7 days') AND owner_user_id = ?
            """,
            (owner,),
        ).fetchone()[0]
        or 0
    )
    latest = _conflict_row(
        conn.execute(
            """
            SELECT conflict_id, table_name, record_sync_id, local_updated_at, remote_updated_at,
                   last_synced_at, resolution, remote_device_id, remote_device_name,
                   detected_at, resolved_at, details_json
            FROM sync_conflicts
            WHERE owner_user_id = ?
            ORDER BY detected_at DESC, id DESC
            LIMIT 1
            """,
            (owner,),
        ).fetchone()
    )
    return {"conflicts_total": total, "conflicts_recent": recent, "latest_conflict": latest}


@app.get("/device/info")
def device_info(service: FinanceService = Depends(get_service)):
    ensure_app_data_initialized()
    with service.db.connect() as conn:
        info = _get_or_create_device_info(conn)
    return {"ok": True, **info, "app_version": app.version}


@app.get("/local-session/context")
def local_session_context(service: FinanceService = Depends(get_service)):
    ensure_app_data_initialized()
    owner = _current_owner()
    with service.db.connect() as conn:
        _ensure_sync_metadata_columns(conn)
        local_counts = {
            table: int(conn.execute(f"SELECT COUNT(*) FROM {table} WHERE owner_user_id = ?", (LOCAL_OWNER_ID,)).fetchone()[0] or 0)
            for table in SYNC_TABLES
        }
        local_claimable_total = sum(local_counts.get(table, 0) for table in CLAIMABLE_LOCAL_DATA_TABLES)
        visible_data = {
            table: int(conn.execute(f"SELECT COUNT(*) FROM {table} WHERE owner_user_id = ?", (owner,)).fetchone()[0] or 0)
            for table in SYNC_TABLES
        }
    return {
        "ok": True,
        "owner_user_id": owner,
        "mode": "local" if owner == LOCAL_OWNER_ID else "cloud",
        "has_local_data": local_claimable_total > 0,
        "has_unassigned_data": local_claimable_total > 0,
        "local_counts": local_counts,
        "local_claimable_total": local_claimable_total,
        "visible_data": visible_data,
    }


@app.post("/local-session/claim-local-data")
async def claim_local_data(request: Request, service: FinanceService = Depends(get_service)):
    payload = await request.json()
    owner = normalize_owner_id(str(payload.get("target_owner_user_id") or payload.get("owner_user_id") or _current_owner()))
    if not owner or owner == LOCAL_OWNER_ID:
        raise HTTPException(status_code=400, detail="Inicia sesion para asociar datos locales a una cuenta.")
    active_cloud_owner = _require_cloud_owner()
    # El destino no puede venir libremente del body: debe coincidir con el owner
    # activo del request para impedir asociaciones accidentales entre cuentas.
    if owner != active_cloud_owner:
        raise HTTPException(status_code=403, detail="La cuenta destino no coincide con la cuenta activa.")
    with service.db.connect() as conn:
        _ensure_sync_metadata_columns(conn)
        summary: dict[str, int] = {table: 0 for table in SYNC_TABLES}
        category_map: dict[int, int] = {}
        tag_map: dict[int, int] = {}
        local_movement_tags = conn.execute(
            """
            SELECT mt.movimiento_id, mt.tag_id
            FROM movimiento_tags mt
            JOIN movimientos m ON m.id = mt.movimiento_id
            WHERE mt.owner_user_id = ? AND m.owner_user_id = ?
            """,
            (LOCAL_OWNER_ID, LOCAL_OWNER_ID),
        ).fetchall()

        category_ids = _local_category_ids_for_claim(conn)
        if category_ids:
            placeholders = ", ".join(["?"] * len(category_ids))
            category_rows = conn.execute(
                f"""
                SELECT id, nombre, tipo
                FROM categorias
                WHERE owner_user_id = ? AND id IN ({placeholders})
                """,
                (LOCAL_OWNER_ID, *category_ids),
            ).fetchall()
            for category in category_rows:
                local_category_id = int(category["id"])
                existing = conn.execute(
                    """
                    SELECT id
                    FROM categorias
                    WHERE owner_user_id = ? AND nombre = ? AND tipo = ?
                    LIMIT 1
                    """,
                    (owner, category["nombre"], category["tipo"]),
                ).fetchone()
                if existing:
                    category_map[local_category_id] = int(existing["id"])
                    continue
                conn.execute(
                    """
                    UPDATE categorias
                    SET owner_user_id = ?,
                        sync_status = 'pending',
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ? AND owner_user_id = ?
                    """,
                    (owner, local_category_id, LOCAL_OWNER_ID),
                )
                category_map[local_category_id] = local_category_id
                summary["categorias"] += 1

        for local_tag_id in sorted({int(row["tag_id"]) for row in local_movement_tags}):
            tag = conn.execute(
                "SELECT id, nombre FROM tags WHERE id = ? AND owner_user_id = ?",
                (local_tag_id, LOCAL_OWNER_ID),
            ).fetchone()
            if not tag:
                continue
            existing = conn.execute(
                "SELECT id FROM tags WHERE owner_user_id = ? AND nombre = ? LIMIT 1",
                (owner, tag["nombre"]),
            ).fetchone()
            if existing:
                tag_map[local_tag_id] = int(existing["id"])
                continue
            conn.execute(
                """
                UPDATE tags
                SET owner_user_id = ?,
                    sync_status = 'pending',
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND owner_user_id = ?
                """,
                (owner, local_tag_id, LOCAL_OWNER_ID),
            )
            tag_map[local_tag_id] = local_tag_id

        def mapped_category_id(value: Any) -> Any:
            if value is None:
                return None
            try:
                return category_map.get(int(value), value)
            except (TypeError, ValueError):
                return value

        def claim_category_table(table: str) -> int:
            rows = conn.execute(
                f"SELECT id, categoria_id FROM {table} WHERE owner_user_id = ?",
                (LOCAL_OWNER_ID,),
            ).fetchall()
            claimed = 0
            for row in rows:
                conn.execute(
                    f"""
                    UPDATE {table}
                    SET owner_user_id = ?,
                        categoria_id = ?,
                        sync_status = 'pending',
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ? AND owner_user_id = ?
                    """,
                    (owner, mapped_category_id(row["categoria_id"]), row["id"], LOCAL_OWNER_ID),
                )
                claimed += 1
            return claimed

        def claim_presupuestos() -> int:
            rows = conn.execute(
                "SELECT id, categoria_id, mes, anio FROM presupuestos WHERE owner_user_id = ?",
                (LOCAL_OWNER_ID,),
            ).fetchall()
            claimed = 0
            for row in rows:
                target_category_id = mapped_category_id(row["categoria_id"])
                existing = conn.execute(
                    """
                    SELECT id
                    FROM presupuestos
                    WHERE owner_user_id = ? AND categoria_id = ? AND mes = ? AND anio = ? AND id != ?
                    LIMIT 1
                    """,
                    (owner, target_category_id, row["mes"], row["anio"], row["id"]),
                ).fetchone()
                if existing:
                    continue
                conn.execute(
                    """
                    UPDATE presupuestos
                    SET owner_user_id = ?,
                        categoria_id = ?,
                        sync_status = 'pending',
                        updated_at = CURRENT_TIMESTAMP
                    WHERE id = ? AND owner_user_id = ?
                    """,
                    (owner, target_category_id, row["id"], LOCAL_OWNER_ID),
                )
                claimed += 1
            return claimed

        summary["movimientos"] = claim_category_table("movimientos")
        for relation in local_movement_tags:
            conn.execute(
                """
                UPDATE movimiento_tags
                SET owner_user_id = ?,
                    tag_id = ?,
                    sync_status = 'pending',
                    updated_at = CURRENT_TIMESTAMP
                WHERE movimiento_id = ? AND tag_id = ? AND owner_user_id = ?
                """,
                (owner, tag_map.get(int(relation["tag_id"]), relation["tag_id"]), relation["movimiento_id"], relation["tag_id"], LOCAL_OWNER_ID),
            )
        summary["gastos_programados"] = claim_category_table("gastos_programados")
        summary["gastos_fijos"] = claim_category_table("gastos_fijos")
        summary["presupuestos"] = claim_presupuestos()
        result = conn.execute(
            """
            UPDATE metas_ahorro
            SET owner_user_id = ?,
                sync_status = 'pending',
                updated_at = CURRENT_TIMESTAMP
            WHERE owner_user_id = ?
            """,
            (owner, LOCAL_OWNER_ID),
        )
        summary["metas_ahorro"] = int(result.rowcount or 0)
        fk_errors = conn.execute("PRAGMA foreign_key_check").fetchall()
        if fk_errors:
            raise HTTPException(status_code=500, detail="La asociacion dejo relaciones invalidas. No se aplicaron cambios.")
    return {
        "ok": True,
        "owner_user_id": owner,
        "claimed": summary,
        "claimed_total": sum(summary.values()),
        "claimable_total": sum(summary.get(table, 0) for table in CLAIMABLE_LOCAL_DATA_TABLES),
    }


@app.get("/sync/overview")
def sync_overview(service: FinanceService = Depends(get_service)):
    ensure_app_data_initialized()
    owner = _current_owner()
    with service.db.connect() as conn:
        _ensure_sync_metadata_columns(conn)
        device = _get_or_create_device_info(conn)
        tables = _sync_status_tables(conn, owner)
        pending_total = sum(item["pending"] for item in tables.values())
        deleted_pending_total = sum(item["deleted_pending"] for item in tables.values())
        return {
            "ok": True,
            "version": app.version,
            "owner_user_id": owner,
            "mode": "local" if owner == LOCAL_OWNER_ID else "cloud",
            **device,
            "has_pending": pending_total > 0,
            "pending_total": pending_total,
            "deleted_pending_total": deleted_pending_total,
            "tables": tables,
            "last_success": _latest_history(conn, "success", owner),
            "last_error": _latest_history(conn, "error", owner),
            "last_remote_change_at": _last_remote_change_at(conn, owner),
            **_conflict_summary(conn, owner),
        }


@app.get("/sync/conflicts")
def sync_conflicts(
    limit: int = Query(default=10, ge=1, le=50),
    offset: int = Query(default=0, ge=0),
    service: FinanceService = Depends(get_service),
):
    ensure_app_data_initialized()
    with service.db.connect() as conn:
        _ensure_sync_history_table(conn)
        rows = conn.execute(
            """
            SELECT conflict_id, table_name, record_sync_id, local_updated_at, remote_updated_at,
                   last_synced_at, resolution, remote_device_id, remote_device_name,
                   detected_at, resolved_at, details_json
            FROM sync_conflicts
            WHERE owner_user_id = ?
            ORDER BY detected_at DESC, id DESC
            LIMIT ? OFFSET ?
            """,
            (_current_owner(), limit, offset),
        ).fetchall()
    return {"ok": True, "items": [_conflict_row(row) for row in rows]}


@app.get("/sync/conflicts/latest")
def sync_conflicts_latest(service: FinanceService = Depends(get_service)):
    ensure_app_data_initialized()
    with service.db.connect() as conn:
        return {"ok": True, **_conflict_summary(conn, _current_owner())}


@app.get("/sync/history")
def sync_history(
    limit: int = Query(default=10, ge=1, le=50),
    offset: int = Query(default=0, ge=0),
    service: FinanceService = Depends(get_service),
):
    ensure_app_data_initialized()
    with service.db.connect() as conn:
        _ensure_sync_history_table(conn)
        rows = conn.execute(
            """
            SELECT sync_id, device_id, mode, status, started_at, finished_at, duration_ms,
                   pending_total, pushed_total, pulled_total, deleted_total,
                   conflicts_total, remote_changes_total, applied_remote_total, kept_local_total,
                   error_message, details_json
            FROM sync_history
            WHERE owner_user_id = ?
            ORDER BY COALESCE(finished_at, started_at) DESC, id DESC
            LIMIT ? OFFSET ?
            """,
            (_current_owner(), limit, offset),
        ).fetchall()
    return {"ok": True, "items": [_history_row(row) for row in rows]}


@app.get("/sync/history/latest")
def sync_history_latest(service: FinanceService = Depends(get_service)):
    ensure_app_data_initialized()
    with service.db.connect() as conn:
        return {
            "ok": True,
            "last_success": _latest_history(conn, "success", _current_owner()),
            "last_error": _latest_history(conn, "error", _current_owner()),
        }


@app.post("/sync/history")
async def sync_history_create(request: Request, service: FinanceService = Depends(get_service)):
    ensure_app_data_initialized()
    payload = await request.json()
    sync_id = str(payload.get("sync_id") or uuid4())
    mode = str(payload.get("mode") or "manual")
    status = str(payload.get("status") or "success")
    if mode not in {"manual", "auto"}:
        raise HTTPException(status_code=400, detail="Modo de sincronizacion invalido.")
    if status not in {"success", "error", "skipped"}:
        raise HTTPException(status_code=400, detail="Estado de sincronizacion invalido.")

    details = payload.get("details") if isinstance(payload.get("details"), dict) else {}
    started_at = str(payload.get("started_at") or datetime.now().isoformat(timespec="seconds"))
    finished_at = payload.get("finished_at")
    with service.db.connect() as conn:
        _ensure_sync_history_table(conn)
        device = _get_or_create_device_info(conn)
        device_id = str(payload.get("device_id") or device["device_id"])
        conn.execute(
            """
            INSERT INTO sync_history (
                sync_id, device_id, owner_user_id, mode, status, started_at, finished_at, duration_ms,
                pending_total, pushed_total, pulled_total, deleted_total,
                conflicts_total, remote_changes_total, applied_remote_total, kept_local_total,
                error_message, details_json
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(sync_id) DO UPDATE SET
                device_id = excluded.device_id,
                owner_user_id = excluded.owner_user_id,
                mode = excluded.mode,
                status = excluded.status,
                started_at = excluded.started_at,
                finished_at = excluded.finished_at,
                duration_ms = excluded.duration_ms,
                pending_total = excluded.pending_total,
                pushed_total = excluded.pushed_total,
                pulled_total = excluded.pulled_total,
                deleted_total = excluded.deleted_total,
                conflicts_total = excluded.conflicts_total,
                remote_changes_total = excluded.remote_changes_total,
                applied_remote_total = excluded.applied_remote_total,
                kept_local_total = excluded.kept_local_total,
                error_message = excluded.error_message,
                details_json = excluded.details_json
            """,
            (
                sync_id,
                device_id,
                _current_owner(),
                mode,
                status,
                started_at,
                str(finished_at) if finished_at else None,
                int(payload.get("duration_ms") or 0),
                int(payload.get("pending_total") or 0),
                int(payload.get("pushed_total") or 0),
                int(payload.get("pulled_total") or 0),
                int(payload.get("deleted_total") or 0),
                int(payload.get("conflicts_total") or 0),
                int(payload.get("remote_changes_total") or 0),
                int(payload.get("applied_remote_total") or 0),
                int(payload.get("kept_local_total") or 0),
                str(payload.get("error_message")) if payload.get("error_message") else None,
                json.dumps(details, ensure_ascii=True),
            ),
        )
    return {"ok": True, "sync_id": sync_id}


def _row_to_dict(row: sqlite3.Row | None) -> dict[str, Any] | None:
    return dict(row) if row is not None else None


def _parse_sync_datetime(value: Any) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        try:
            return datetime.strptime(str(value), "%Y-%m-%d %H:%M:%S")
        except ValueError:
            return None


def _remote_is_newer(item: dict[str, Any], local: sqlite3.Row) -> bool:
    """Compare cloud revisions, not notebook clocks. Device id is a deterministic tie-breaker."""
    remote_dt = _parse_sync_datetime(item.get("remote_updated_at"))
    local_dt = _parse_sync_datetime(local["last_remote_updated_at"])
    if remote_dt is None:
        return False
    if local_dt is None or remote_dt > local_dt:
        return True
    if remote_dt < local_dt:
        return False
    remote_device = str(item.get("last_modified_device_id") or "")
    local_device = str(local["last_remote_device_id"] or "")
    return remote_device > local_device


def _changed_after(value: Any, baseline: Any) -> bool:
    value_dt = _parse_sync_datetime(value)
    baseline_dt = _parse_sync_datetime(baseline)
    if value_dt is None or baseline_dt is None:
        return False
    return value_dt > baseline_dt


def _basic_conflict_detected(local: sqlite3.Row, item: dict[str, Any], current_device_id: str | None = None) -> bool:
    if str(local["sync_status"] or "") != "pending":
        return False
    remote_device_id = str(item.get("last_modified_device_id") or "").strip()
    if current_device_id and remote_device_id and remote_device_id == current_device_id:
        return False
    remote_revision = str(item.get("remote_updated_at") or "").strip()
    local_revision = str(local["last_remote_updated_at"] or "").strip()
    if remote_revision:
        return remote_revision != local_revision
    last_synced_at = local["last_synced_at"]
    return _changed_after(local["updated_at"], last_synced_at) and _changed_after(item.get("updated_at"), last_synced_at)


def _record_sync_conflict(
    conn: sqlite3.Connection,
    table: str,
    local: sqlite3.Row,
    item: dict[str, Any],
    resolution: str,
    now: str,
    result: dict[str, Any],
) -> None:
    conflict_id = str(uuid4())
    remote_device_id = item.get("last_modified_device_id") or item.get("remote_device_id")
    remote_device_name = item.get("last_modified_device_name") or item.get("remote_device_name")
    conn.execute(
        """
        INSERT INTO sync_conflicts (
            conflict_id, owner_user_id, table_name, record_sync_id, local_updated_at, remote_updated_at,
            last_synced_at, resolution, remote_device_id, remote_device_name,
            detected_at, resolved_at, details_json
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            conflict_id,
            _current_owner(),
            table,
            item.get("sync_id"),
            local["updated_at"],
            item.get("updated_at"),
            local["last_synced_at"],
            resolution,
            remote_device_id,
            remote_device_name,
            now,
            now,
            json.dumps(
                {
                    "message": "Cambios detectados en mas de un dispositivo. Se aplico last-write-wins.",
                    "local_sync_status": local["sync_status"],
                },
                ensure_ascii=True,
            ),
        ),
    )
    result["conflicts_total"] = int(result.get("conflicts_total", 0)) + 1
    result.setdefault("conflicts_by_table", {})
    result["conflicts_by_table"][table] = int(result["conflicts_by_table"].get(table, 0)) + 1


def _remote_meta_values(item: dict[str, Any]) -> tuple[Any, Any, Any]:
    return (
        item.get("last_modified_device_id"),
        item.get("last_modified_device_name"),
        item.get("last_modified_at") or item.get("updated_at"),
    )


SYNC_TABLES = ("categorias", "tags", "metas_ahorro", "gastos_programados", "gastos_fijos", "presupuestos", "movimientos", "movimiento_tags")
INTEGRITY_SYNC_TABLES = SYNC_TABLES
CLAIMABLE_LOCAL_DATA_TABLES = ("movimientos", "metas_ahorro", "gastos_programados", "gastos_fijos", "presupuestos")
EXPECTED_LOCAL_COLUMNS: dict[str, set[str]] = {
    "categorias": {"id", "nombre", "tipo", "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status"},
    "movimientos": {"id", "fecha", "tipo", "categoria_id", "monto", "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status"},
    "metas_ahorro": {"id", "nombre", "monto_objetivo", "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status"},
    "gastos_programados": {"id", "descripcion", "categoria_id", "monto_estimado", "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status"},
    "gastos_fijos": {"id", "categoria_id", "monto", "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status"},
    "presupuestos": {"id", "categoria_id", "mes", "anio", "monto", "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status"},
    "tags": {"id", "nombre", "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status"},
    "movimiento_tags": {"movimiento_id", "tag_id", "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status"},
}


def _db_integrity_issue(code: str, severity: str, table: str | None, count: int, repairable: bool) -> dict[str, Any]:
    return {"code": code, "severity": severity, "table": table, "count": count, "repairable": repairable}


def _open_local_db(read_only: bool = False) -> sqlite3.Connection:
    if read_only:
        conn = sqlite3.connect(f"{WEB_DB_PATH.resolve().as_uri()}?mode=ro", uri=True, timeout=30)
    else:
        conn = sqlite3.connect(WEB_DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout = 30000")
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def _table_columns(conn: sqlite3.Connection, table: str) -> set[str]:
    return {str(row["name"]) for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}


def _count_rows(conn: sqlite3.Connection, query: str, params: tuple[Any, ...] = ()) -> int:
    return int(conn.execute(query, params).fetchone()[0] or 0)


def _append_count_issue(
    issues: list[dict[str, Any]],
    conn: sqlite3.Connection,
    code: str,
    severity: str,
    table: str,
    query: str,
    repairable: bool = False,
) -> None:
    count = _count_rows(conn, query)
    if count:
        issues.append(_db_integrity_issue(code, severity, table, count, repairable))


def _has_index_with_columns(conn: sqlite3.Connection, table: str, expected: tuple[str, ...], unique: bool | None = None) -> bool:
    for index in conn.execute(f"PRAGMA index_list({table})").fetchall():
        if unique is not None and bool(index["unique"]) != unique:
            continue
        columns = tuple(str(row["name"]) for row in conn.execute(f"PRAGMA index_info({index['name']})").fetchall())
        if columns == expected:
            return True
    return False


def _tables_have_columns(conn: sqlite3.Connection, requirements: dict[str, set[str]], existing_tables: set[str]) -> bool:
    return all(table in existing_tables and columns.issubset(_table_columns(conn, table)) for table, columns in requirements.items())


def _local_db_integrity_report() -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    detected_schema_version: str | None = None
    if not WEB_DB_PATH.exists():
        issues.append(_db_integrity_issue("database_missing", "critical", None, 1, False))
        return _build_local_db_integrity_response(issues, detected_schema_version)

    try:
        with closing(_open_local_db(read_only=True)) as conn:
            table_rows = conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'").fetchall()
            existing_tables = {str(row["name"]) for row in table_rows}
            for table, required_columns in EXPECTED_LOCAL_COLUMNS.items():
                if table not in existing_tables:
                    issues.append(_db_integrity_issue("missing_table", "critical", table, 1, False))
                    continue
                columns = _table_columns(conn, table)
                for column in sorted(required_columns - columns):
                    issues.append(_db_integrity_issue(f"missing_column_{column}", "critical", table, 1, False))

            fk_issues = conn.execute("PRAGMA foreign_key_check").fetchall()
            if fk_issues:
                issues.append(_db_integrity_issue("foreign_key_violation", "critical", None, len(fk_issues), False))

            integrity_row = conn.execute("PRAGMA integrity_check").fetchone()
            if not integrity_row or str(integrity_row[0]).lower() != "ok":
                issues.append(_db_integrity_issue("sqlite_integrity_check_failed", "critical", None, 1, False))

            for table in INTEGRITY_SYNC_TABLES:
                if table not in existing_tables:
                    continue
                columns = _table_columns(conn, table)
                if "owner_user_id" in columns:
                    count = int(conn.execute(f"SELECT COUNT(*) FROM {table} WHERE owner_user_id IS NULL OR trim(owner_user_id) = ''").fetchone()[0] or 0)
                    if count:
                        issues.append(_db_integrity_issue("missing_owner_user_id", "critical", table, count, False))
                for column in ("created_at", "updated_at"):
                    if column not in columns:
                        continue
                    count = int(conn.execute(f"SELECT COUNT(*) FROM {table} WHERE {column} IS NULL OR trim({column}) = '' OR datetime({column}) IS NULL").fetchone()[0] or 0)
                    if count:
                        issues.append(_db_integrity_issue(f"invalid_{column}", "warning", table, count, True))
                if "deleted_at" in columns:
                    count = int(conn.execute(f"SELECT COUNT(*) FROM {table} WHERE deleted_at IS NOT NULL AND trim(deleted_at) <> '' AND datetime(deleted_at) IS NULL").fetchone()[0] or 0)
                    if count:
                        issues.append(_db_integrity_issue("invalid_deleted_at", "warning", table, count, False))
                if "sync_status" in columns:
                    count = int(conn.execute(f"SELECT COUNT(*) FROM {table} WHERE sync_status IS NULL OR trim(sync_status) = '' OR sync_status NOT IN ('pending', 'synced')").fetchone()[0] or 0)
                    if count:
                        issues.append(_db_integrity_issue("invalid_sync_status", "warning", table, count, True))
                if "sync_id" in columns:
                    count = int(conn.execute(f"SELECT COUNT(*) FROM {table} WHERE sync_id IS NULL OR trim(sync_id) = ''").fetchone()[0] or 0)
                    if count:
                        issues.append(_db_integrity_issue("missing_sync_id", "warning", table, count, True))
                    duplicates = conn.execute(
                        f"SELECT COUNT(*) FROM (SELECT sync_id FROM {table} WHERE sync_id IS NOT NULL AND trim(sync_id) <> '' GROUP BY sync_id HAVING COUNT(*) > 1)"
                    ).fetchone()
                    duplicate_count = int(duplicates[0] or 0)
                    if duplicate_count:
                        issues.append(_db_integrity_issue("duplicate_sync_id", "critical", table, duplicate_count, False))

            relationship_checks = [
                (
                    "cross_owner_category_reference",
                    "movimientos",
                    {"movimientos": {"categoria_id", "owner_user_id"}, "categorias": {"id", "owner_user_id"}},
                    "SELECT COUNT(*) FROM movimientos child JOIN categorias parent ON parent.id = child.categoria_id WHERE child.owner_user_id <> parent.owner_user_id",
                ),
                (
                    "cross_owner_meta_reference",
                    "movimientos",
                    {"movimientos": {"meta_id", "owner_user_id"}, "metas_ahorro": {"id", "owner_user_id"}},
                    "SELECT COUNT(*) FROM movimientos child JOIN metas_ahorro parent ON parent.id = child.meta_id WHERE child.meta_id IS NOT NULL AND child.owner_user_id <> parent.owner_user_id",
                ),
                (
                    "cross_owner_category_reference",
                    "gastos_fijos",
                    {"gastos_fijos": {"categoria_id", "owner_user_id"}, "categorias": {"id", "owner_user_id"}},
                    "SELECT COUNT(*) FROM gastos_fijos child JOIN categorias parent ON parent.id = child.categoria_id WHERE child.owner_user_id <> parent.owner_user_id",
                ),
                (
                    "cross_owner_category_reference",
                    "gastos_programados",
                    {"gastos_programados": {"categoria_id", "owner_user_id"}, "categorias": {"id", "owner_user_id"}},
                    "SELECT COUNT(*) FROM gastos_programados child JOIN categorias parent ON parent.id = child.categoria_id WHERE child.owner_user_id <> parent.owner_user_id",
                ),
                (
                    "cross_owner_category_reference",
                    "presupuestos",
                    {"presupuestos": {"categoria_id", "owner_user_id"}, "categorias": {"id", "owner_user_id"}},
                    "SELECT COUNT(*) FROM presupuestos child JOIN categorias parent ON parent.id = child.categoria_id WHERE child.owner_user_id <> parent.owner_user_id",
                ),
                (
                    "cross_owner_movement_tag_reference",
                    "movimiento_tags",
                    {"movimiento_tags": {"movimiento_id", "owner_user_id"}, "movimientos": {"id", "owner_user_id"}},
                    "SELECT COUNT(*) FROM movimiento_tags child JOIN movimientos parent ON parent.id = child.movimiento_id WHERE child.owner_user_id <> parent.owner_user_id",
                ),
                (
                    "cross_owner_tag_reference",
                    "movimiento_tags",
                    {"movimiento_tags": {"tag_id", "owner_user_id"}, "tags": {"id", "owner_user_id"}},
                    "SELECT COUNT(*) FROM movimiento_tags child JOIN tags parent ON parent.id = child.tag_id WHERE child.owner_user_id <> parent.owner_user_id",
                ),
            ]
            for code, table, requirements, query in relationship_checks:
                if _tables_have_columns(conn, requirements, existing_tables):
                    _append_count_issue(issues, conn, code, "critical", table, query)

            logical_uniques = [
                ("duplicate_owner_category", "categorias", ("owner_user_id", "nombre", "tipo")),
                ("duplicate_owner_tag", "tags", ("owner_user_id", "nombre")),
                ("duplicate_owner_budget", "presupuestos", ("owner_user_id", "categoria_id", "mes", "anio")),
                ("duplicate_owner_movement_tag", "movimiento_tags", ("owner_user_id", "movimiento_id", "tag_id")),
            ]
            for code, table, columns in logical_uniques:
                if table not in existing_tables or not set(columns).issubset(_table_columns(conn, table)):
                    continue
                grouped = ", ".join(columns)
                _append_count_issue(
                    issues,
                    conn,
                    code,
                    "critical",
                    table,
                    f"SELECT COUNT(*) FROM (SELECT 1 FROM {table} GROUP BY {grouped} HAVING COUNT(*) > 1)",
                )

            owner_index_tables = (*INTEGRITY_SYNC_TABLES, "sync_history", "sync_conflicts")
            for table in owner_index_tables:
                if table in existing_tables and "owner_user_id" in _table_columns(conn, table):
                    has_owner_index = any(
                        "owner_user_id" in tuple(
                            str(row["name"]) for row in conn.execute(f"PRAGMA index_info({index['name']})").fetchall()
                        )
                        for index in conn.execute(f"PRAGMA index_list({table})").fetchall()
                    )
                    if not has_owner_index:
                        issues.append(_db_integrity_issue("missing_owner_index", "warning", table, 1, False))

            expected_unique_indexes = [
                ("categorias", ("owner_user_id", "nombre", "tipo")),
                ("tags", ("owner_user_id", "nombre")),
                ("presupuestos", ("owner_user_id", "categoria_id", "mes", "anio")),
            ]
            for table, columns in expected_unique_indexes:
                if table in existing_tables and not _has_index_with_columns(conn, table, columns, unique=True):
                    issues.append(_db_integrity_issue("missing_owner_unique_constraint", "critical", table, 1, False))

            # Estos CHECK forman parte de la semantica financiera minima. Detectar
            # su ausencia evita operar sobre DBs parcialmente reconstruidas.
            critical_sql_fragments = {
                "categorias": ("check(tipo in ('ingreso', 'gasto', 'ahorro', 'inversion'))",),
                "movimientos": ("check(tipo in ('ingreso', 'gasto', 'ahorro', 'inversion'))", "check(monto >= 0)"),
                "presupuestos": ("check(mes between 1 and 12)", "check(monto > 0)"),
            }
            for table, fragments in critical_sql_fragments.items():
                if table not in existing_tables:
                    continue
                create_sql_row = conn.execute("SELECT sql FROM sqlite_master WHERE type = 'table' AND name = ?", (table,)).fetchone()
                normalized_sql = "".join(str(create_sql_row["sql"] or "").lower().split()) if create_sql_row else ""
                for fragment in fragments:
                    if "".join(fragment.split()) not in normalized_sql:
                        issues.append(_db_integrity_issue("missing_critical_constraint", "critical", table, 1, False))

            if "app_config" in existing_tables:
                invalid_cursors = conn.execute(
                    "SELECT key, value FROM app_config WHERE key LIKE 'sync_pull_cursor:%'"
                ).fetchall()
                invalid_cursor_count = sum(1 for row in invalid_cursors if _parse_sync_datetime(row["value"]) is None)
                if invalid_cursor_count:
                    issues.append(_db_integrity_issue("invalid_sync_cursor", "warning", "app_config", invalid_cursor_count, True))

            date_checks = [
                ("movimientos", "fecha"),
                ("gastos_programados", "fecha_vencimiento"),
                ("metas_ahorro", "fecha_objetivo"),
            ]
            for table, column in date_checks:
                if table in existing_tables and column in _table_columns(conn, table):
                    _append_count_issue(
                        issues,
                        conn,
                        f"invalid_{column}",
                        "warning",
                        table,
                        f"SELECT COUNT(*) FROM {table} WHERE {column} IS NOT NULL AND trim({column}) <> '' AND date({column}) IS NULL",
                    )

            for table in ("sync_history", "sync_conflicts"):
                if table in existing_tables and "owner_user_id" in _table_columns(conn, table):
                    _append_count_issue(
                        issues,
                        conn,
                        "missing_owner_user_id",
                        "critical",
                        table,
                        f"SELECT COUNT(*) FROM {table} WHERE owner_user_id IS NULL OR trim(owner_user_id) = ''",
                    )

            if "app_config" not in existing_tables:
                issues.append(_db_integrity_issue("schema_version_unavailable", "warning", "app_config", 1, False))
            else:
                schema_version = conn.execute("SELECT value FROM app_config WHERE key = 'schema_version'").fetchone()
                if not schema_version or not str(schema_version["value"] or "").strip():
                    issues.append(_db_integrity_issue("schema_version_missing", "warning", "app_config", 1, False))
                else:
                    detected_schema_version = str(schema_version["value"]).strip()
                    if detected_schema_version != CURRENT_SCHEMA_VERSION:
                        issues.append(_db_integrity_issue("schema_version_unknown", "critical", "app_config", 1, False))
    except sqlite3.Error as exc:
        _logger.exception("Fallo chequeo de integridad local. error_type=%s", type(exc).__name__)
        issues.append(_db_integrity_issue("database_unreadable", "critical", None, 1, False))

    return _build_local_db_integrity_response(issues, detected_schema_version)


def _build_local_db_integrity_response(issues: list[dict[str, Any]], detected_schema_version: str | None = None) -> dict[str, Any]:
    critical_count = sum(int(issue["count"]) for issue in issues if issue["severity"] == "critical")
    warning_count = sum(int(issue["count"]) for issue in issues if issue["severity"] == "warning")
    repairable_count = sum(int(issue["count"]) for issue in issues if issue["repairable"])
    status = "critical" if critical_count else "warning" if warning_count else "healthy"
    summary = ["No se encontraron problemas."] if status == "healthy" else [
        "Encontramos datos locales que necesitan una revision." if status == "warning" else "Encontramos problemas locales que requieren reparacion antes de sincronizar."
    ]
    return {
        "ok": status != "critical",
        "status": status,
        "issues_count": critical_count + warning_count,
        "warnings_count": warning_count,
        "repairable_count": repairable_count,
        "backup_recommended": status != "healthy",
        "schema_version": detected_schema_version,
        "expected_schema_version": CURRENT_SCHEMA_VERSION,
        "safe_summary": summary,
        "issues": issues,
    }


def _create_local_backup() -> Path:
    if not WEB_DB_PATH.exists():
        raise HTTPException(status_code=404, detail="No existe una base de datos local para respaldar.")
    backup_dir = get_backup_dir()
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = backup_dir / f"finanzas_backup_{stamp}.db"
    if target.exists():
        target = backup_dir / f"finanzas_backup_{stamp}_{uuid4().hex[:8]}.db"
    with closing(_open_local_db(read_only=True)) as source, closing(sqlite3.connect(target)) as destination:
        source.backup(destination)
    return target


@app.get("/local/db-integrity")
def local_db_integrity():
    report = _local_db_integrity_report()
    _logger.info("[db-integrity] status=%s issue_codes=%s", report["status"], [issue["code"] for issue in report["issues"]])
    return report


@app.post("/local/backup")
def local_backup():
    created = _create_local_backup()
    _logger.info("[db-backup] created=%s", created.name)
    return {"ok": True, "backup_path": str(created), "created_at": datetime.now().isoformat(timespec="seconds")}


@app.post("/local/db-repair")
def local_db_repair():
    backup = _create_local_backup()
    before = _local_db_integrity_report()
    _logger.info("[db-repair] start issue_codes=%s", [issue["code"] for issue in before["issues"]])
    if any(issue["severity"] == "critical" for issue in before["issues"]):
        return JSONResponse(
            status_code=409,
            content={
                "ok": False,
                "backup_created": True,
                "repaired_count": 0,
                "unresolved_count": before["issues_count"],
                "safe_summary": ["Se creo un backup antes de revisar.", "Encontramos problemas que requieren revision manual.", "Tus datos no fueron modificados."],
            },
        )

    repaired_count = 0
    now = datetime.now().isoformat(timespec="seconds")
    with closing(_open_local_db()) as conn:
        with conn:
            for table in INTEGRITY_SYNC_TABLES:
                columns = _table_columns(conn, table)
                if "created_at" in columns:
                    cursor = conn.execute(f"UPDATE {table} SET created_at = ? WHERE created_at IS NULL OR trim(created_at) = '' OR datetime(created_at) IS NULL", (now,))
                    repaired_count += max(cursor.rowcount, 0)
                if "updated_at" in columns:
                    cursor = conn.execute(f"UPDATE {table} SET updated_at = COALESCE(NULLIF(created_at, ''), ?) WHERE updated_at IS NULL OR trim(updated_at) = '' OR datetime(updated_at) IS NULL", (now,))
                    repaired_count += max(cursor.rowcount, 0)
                if "sync_status" in columns:
                    cursor = conn.execute(f"UPDATE {table} SET sync_status = 'pending' WHERE sync_status IS NULL OR trim(sync_status) = '' OR sync_status NOT IN ('pending', 'synced')")
                    repaired_count += max(cursor.rowcount, 0)
                if "sync_id" in columns:
                    rows = conn.execute(f"SELECT rowid FROM {table} WHERE sync_id IS NULL OR trim(sync_id) = ''").fetchall()
                    for row in rows:
                        conn.execute(f"UPDATE {table} SET sync_id = ? WHERE rowid = ?", (str(uuid4()), row["rowid"]))
                    repaired_count += len(rows)
            # Borrar un cursor corrupto es seguro: la proxima sync realiza un
            # full-pull y reconstruye el cursor sin eliminar datos financieros.
            existing_tables = {str(row["name"]) for row in conn.execute("SELECT name FROM sqlite_master WHERE type = 'table'").fetchall()}
            if "app_config" in existing_tables:
                cursor_rows = conn.execute("SELECT key, value FROM app_config WHERE key LIKE 'sync_pull_cursor:%'").fetchall()
                for row in cursor_rows:
                    if _parse_sync_datetime(row["value"]) is None:
                        conn.execute("DELETE FROM app_config WHERE key = ?", (row["key"],))
                        repaired_count += 1

    after = _local_db_integrity_report()
    unresolved_count = int(after["issues_count"])
    _logger.info("[db-repair] finish repaired=%s unresolved=%s issue_codes=%s", repaired_count, unresolved_count, [issue["code"] for issue in after["issues"]])
    return {
        "ok": after["status"] != "critical",
        "backup_created": True,
        "backup_path": str(backup),
        "repaired_count": repaired_count,
        "unresolved_count": unresolved_count,
        "safe_summary": ["Se creo un backup antes de reparar.", "No se eliminaron datos financieros."] + list(after["safe_summary"]),
    }


def _local_category_ids_for_claim(conn: sqlite3.Connection) -> list[int]:
    category_ids: set[int] = set()
    queries = [
        "SELECT categoria_id FROM movimientos WHERE owner_user_id = ? AND categoria_id IS NOT NULL",
        "SELECT categoria_id FROM gastos_programados WHERE owner_user_id = ? AND categoria_id IS NOT NULL",
        "SELECT categoria_id FROM gastos_fijos WHERE owner_user_id = ? AND categoria_id IS NOT NULL",
        "SELECT categoria_id FROM presupuestos WHERE owner_user_id = ? AND categoria_id IS NOT NULL",
    ]
    for query in queries:
        for row in conn.execute(query, (LOCAL_OWNER_ID,)).fetchall():
            if row["categoria_id"] is not None:
                category_ids.add(int(row["categoria_id"]))
    return sorted(category_ids)


SYNC_SELECTS: dict[str, str] = {
    "categorias": """
        SELECT sync_id, nombre, tipo, owner_user_id, created_at, updated_at, deleted_at, sync_status, last_synced_at,
               last_remote_updated_at
        FROM categorias
    """,
    "movimientos": """
        SELECT
          m.sync_id, m.fecha, m.tipo, m.monto, m.descripcion, m.categoria_id,
          c.sync_id AS categoria_sync_id,
          m.owner_user_id, m.created_at, m.updated_at, m.deleted_at, m.sync_status, m.last_synced_at,
          m.last_remote_updated_at
        FROM movimientos m
        LEFT JOIN categorias c ON c.id = m.categoria_id AND c.owner_user_id = m.owner_user_id
    """,
    "metas_ahorro": """
        SELECT sync_id, nombre, monto_objetivo, monto_inicial, fecha_objetivo, descripcion, estado,
               owner_user_id, created_at, updated_at, deleted_at, sync_status, last_synced_at, last_remote_updated_at
        FROM metas_ahorro
    """,
    "gastos_programados": """
        SELECT gp.sync_id, gp.descripcion, gp.monto_estimado, gp.fecha_vencimiento, gp.estado,
               gp.es_recurrente, gp.frecuencia, gp.categoria_id, c.sync_id AS categoria_sync_id,
               gp.owner_user_id, gp.created_at, gp.updated_at, gp.deleted_at, gp.sync_status, gp.last_synced_at,
               gp.last_remote_updated_at
        FROM gastos_programados gp
        LEFT JOIN categorias c ON c.id = gp.categoria_id AND c.owner_user_id = gp.owner_user_id
    """,
    "gastos_fijos": """
        SELECT gf.sync_id, gf.descripcion, gf.monto, gf.dia_vencimiento, gf.activo,
               gf.categoria_id, c.sync_id AS categoria_sync_id,
               gf.owner_user_id, gf.created_at, gf.updated_at, gf.deleted_at, gf.sync_status, gf.last_synced_at,
               gf.last_remote_updated_at
        FROM gastos_fijos gf
        LEFT JOIN categorias c ON c.id = gf.categoria_id AND c.owner_user_id = gf.owner_user_id
    """,
    "presupuestos": """
        SELECT p.sync_id, p.mes, p.anio, p.monto, p.categoria_id, c.sync_id AS categoria_sync_id,
               p.owner_user_id, p.created_at, p.updated_at, p.deleted_at, p.sync_status, p.last_synced_at,
               p.last_remote_updated_at
        FROM presupuestos p
        LEFT JOIN categorias c ON c.id = p.categoria_id AND c.owner_user_id = p.owner_user_id
    """,
    "tags": """
        SELECT sync_id, nombre, color, owner_user_id, created_at, updated_at, deleted_at, sync_status, last_synced_at,
               last_remote_updated_at
        FROM tags
    """,
    "movimiento_tags": """
        SELECT mt.sync_id, m.sync_id AS movimiento_sync_id, t.sync_id AS tag_sync_id,
               mt.owner_user_id, mt.created_at, mt.updated_at, mt.deleted_at, mt.sync_status, mt.last_synced_at,
               mt.last_remote_updated_at
        FROM movimiento_tags mt
        JOIN movimientos m ON m.id = mt.movimiento_id AND m.owner_user_id = mt.owner_user_id
        JOIN tags t ON t.id = mt.tag_id AND t.owner_user_id = mt.owner_user_id
    """,
}

SYNC_STATUS_PREFIX: dict[str, str] = {
    "categorias": "",
    "movimientos": "m.",
    "metas_ahorro": "",
    "gastos_programados": "gp.",
    "gastos_fijos": "gf.",
    "presupuestos": "p.",
    "tags": "",
    "movimiento_tags": "mt.",
}


def _sync_pending_rows(conn: sqlite3.Connection, table: str, owner: str) -> list[dict[str, Any]]:
    prefix = SYNC_STATUS_PREFIX[table]
    query = f"""
    {SYNC_SELECTS[table]}
    WHERE ({prefix}sync_status = 'pending' OR {prefix}sync_status IS NULL OR trim({prefix}sync_status) = '')
      AND {prefix}owner_user_id = ?
    """
    return [dict(row) for row in conn.execute(query, (owner,)).fetchall()]


def _local_category_id(conn: sqlite3.Connection, item: dict[str, Any]) -> int | None:
    categoria_sync_id = str(item.get("categoria_sync_id") or "").strip()
    if categoria_sync_id:
        row = conn.execute("SELECT id FROM categorias WHERE sync_id = ? AND owner_user_id = ?", (categoria_sync_id, _current_owner())).fetchone()
        if row:
            return int(row["id"])
    categoria_id = item.get("categoria_id")
    if categoria_id is not None:
        row = conn.execute("SELECT id FROM categorias WHERE id = ? AND owner_user_id = ?", (categoria_id, _current_owner())).fetchone()
        if row:
            return int(row["id"])
    return None


def _local_sync_id_record_id(conn: sqlite3.Connection, table: str, sync_id: Any) -> int | None:
    clean_sync_id = str(sync_id or "").strip()
    if not clean_sync_id:
        return None
    row = conn.execute(
        f"SELECT id FROM {table} WHERE sync_id = ? AND owner_user_id = ?",
        (clean_sync_id, _current_owner()),
    ).fetchone()
    return int(row["id"]) if row else None


def _sync_status_value(item: dict[str, Any]) -> str:
    return "synced"


def _apply_simple_remote(
    conn: sqlite3.Connection,
    table: str,
    item: dict[str, Any],
    fields: list[str],
    result: dict[str, int],
    now: str,
    current_device_id: str | None = None,
) -> None:
    sync_id = str(item.get("sync_id") or "").strip()
    owner = _current_owner()
    if not sync_id:
        result[f"{table}_skipped"] = result.get(f"{table}_skipped", 0) + 1
        return
    local = conn.execute(f"SELECT * FROM {table} WHERE sync_id = ? AND owner_user_id = ?", (sync_id, owner)).fetchone()
    values = [item.get(field) for field in fields]
    if local:
        conflict = _basic_conflict_detected(local, item, current_device_id)
        if conflict and not _remote_is_newer(item, local):
            _record_sync_conflict(conn, table, local, item, "kept_local", now, result)
            result[f"{table}_kept_local"] = result.get(f"{table}_kept_local", 0) + 1
            result["kept_local_total"] = int(result.get("kept_local_total", 0)) + 1
            return
        if _remote_is_newer(item, local):
            if conflict:
                _record_sync_conflict(conn, table, local, item, "applied_remote", now, result)
            assignments = ", ".join([f"{field} = ?" for field in fields])
            conn.execute(
                f"""
                UPDATE {table}
                SET {assignments},
                    updated_at = ?, deleted_at = ?, sync_status = ?, last_synced_at = ?,
                    last_remote_device_id = ?, last_remote_device_name = ?, last_remote_updated_at = ?
                WHERE sync_id = ? AND owner_user_id = ?
                """,
                (
                    *values,
                    item.get("updated_at") or now,
                    item.get("deleted_at"),
                    _sync_status_value(item),
                    now,
                    *_remote_meta_values(item),
                    sync_id,
                    owner,
                ),
            )
            result[f"{table}_updated"] = result.get(f"{table}_updated", 0) + 1
            result["applied_remote_total"] = int(result.get("applied_remote_total", 0)) + 1
        return

    columns = ", ".join([*fields, "owner_user_id", "sync_id", "created_at", "updated_at", "deleted_at", "sync_status", "last_synced_at", "last_remote_device_id", "last_remote_device_name", "last_remote_updated_at"])
    placeholders = ", ".join(["?"] * (len(fields) + 10))
    conn.execute(
        f"INSERT INTO {table} ({columns}) VALUES ({placeholders})",
        (
            *values,
            owner,
            sync_id,
            item.get("created_at") or now,
            item.get("updated_at") or now,
            item.get("deleted_at"),
            _sync_status_value(item),
            now,
            *_remote_meta_values(item),
        ),
    )
    result[f"{table}_inserted"] = result.get(f"{table}_inserted", 0) + 1
    result["applied_remote_total"] = int(result.get("applied_remote_total", 0)) + 1


def _apply_remote_tag(
    conn: sqlite3.Connection,
    item: dict[str, Any],
    result: dict[str, Any],
    now: str,
    current_device_id: str | None,
) -> None:
    sync_id = str(item.get("sync_id") or "").strip()
    nombre = str(item.get("nombre") or "").strip()
    owner = _current_owner()
    if not sync_id or not nombre:
        result["tags_skipped"] += 1
        return
    local = conn.execute("SELECT * FROM tags WHERE sync_id = ? AND owner_user_id = ?", (sync_id, owner)).fetchone()
    if local:
        _apply_simple_remote(conn, "tags", item, ["nombre", "color"], result, now, current_device_id)
        return
    existing = conn.execute("SELECT * FROM tags WHERE nombre = ? AND owner_user_id = ?", (nombre, owner)).fetchone()
    if existing:
        if _remote_is_newer(item, existing):
            conn.execute(
                """
                UPDATE tags
                SET sync_id = ?, color = ?, updated_at = ?, deleted_at = ?, sync_status = 'synced', last_synced_at = ?,
                    last_remote_device_id = ?, last_remote_device_name = ?, last_remote_updated_at = ?
                WHERE id = ? AND owner_user_id = ?
                """,
                (sync_id, item.get("color"), item.get("updated_at") or now, item.get("deleted_at"), now, *_remote_meta_values(item), existing["id"], owner),
            )
            result["tags_updated"] += 1
            result["applied_remote_total"] += 1
        return
    _apply_simple_remote(conn, "tags", item, ["nombre", "color"], result, now, current_device_id)


def _apply_remote_movement_tag(conn: sqlite3.Connection, item: dict[str, Any], result: dict[str, Any], now: str) -> None:
    sync_id = str(item.get("sync_id") or "").strip()
    owner = _current_owner()
    movimiento_id = _local_sync_id_record_id(conn, "movimientos", item.get("movimiento_sync_id"))
    tag_id = _local_sync_id_record_id(conn, "tags", item.get("tag_sync_id"))
    if not sync_id or movimiento_id is None or tag_id is None:
        result["movimiento_tags_skipped"] += 1
        return
    local = conn.execute("SELECT * FROM movimiento_tags WHERE sync_id = ? AND owner_user_id = ?", (sync_id, owner)).fetchone()
    if local:
        if _remote_is_newer(item, local):
            conn.execute(
                """
                UPDATE movimiento_tags
                SET movimiento_id = ?, tag_id = ?, updated_at = ?, deleted_at = ?, sync_status = 'synced',
                    last_synced_at = ?, last_remote_device_id = ?, last_remote_device_name = ?, last_remote_updated_at = ?
                WHERE sync_id = ? AND owner_user_id = ?
                """,
                (movimiento_id, tag_id, item.get("updated_at") or now, item.get("deleted_at"), now, *_remote_meta_values(item), sync_id, owner),
            )
            result["movimiento_tags_updated"] += 1
            result["applied_remote_total"] += 1
        return
    existing = conn.execute(
        "SELECT * FROM movimiento_tags WHERE movimiento_id = ? AND tag_id = ? AND owner_user_id = ?",
        (movimiento_id, tag_id, owner),
    ).fetchone()
    if existing:
        if _remote_is_newer(item, existing):
            conn.execute(
                """
                UPDATE movimiento_tags
                SET sync_id = ?, updated_at = ?, deleted_at = ?, sync_status = 'synced', last_synced_at = ?,
                    last_remote_device_id = ?, last_remote_device_name = ?, last_remote_updated_at = ?
                WHERE movimiento_id = ? AND tag_id = ? AND owner_user_id = ?
                """,
                (sync_id, item.get("updated_at") or now, item.get("deleted_at"), now, *_remote_meta_values(item), movimiento_id, tag_id, owner),
            )
            result["movimiento_tags_updated"] += 1
            result["applied_remote_total"] += 1
        return
    conn.execute(
        """
        INSERT INTO movimiento_tags (
            movimiento_id, tag_id, owner_user_id, sync_id, created_at, updated_at, deleted_at, sync_status,
            last_synced_at, last_remote_device_id, last_remote_device_name, last_remote_updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, 'synced', ?, ?, ?, ?)
        """,
        (movimiento_id, tag_id, owner, sync_id, item.get("created_at") or now, item.get("updated_at") or now, item.get("deleted_at"), now, *_remote_meta_values(item)),
    )
    result["movimiento_tags_inserted"] += 1
    result["applied_remote_total"] += 1


@app.get("/sync/pending")
def sync_pending(service: FinanceService = Depends(get_service)):
    owner = _require_cloud_owner()
    ensure_app_data_initialized()
    with service.db.connect() as conn:
        _ensure_sync_metadata_columns(conn)
        payload = {table: _sync_pending_rows(conn, table, owner) for table in SYNC_TABLES}
    return {"ok": True, **payload}


@app.get("/sync/cursor")
def sync_cursor_get(service: FinanceService = Depends(get_service)):
    owner = _require_cloud_owner()
    with service.db.connect() as conn:
        return {"ok": True, "cursor": _config_get(conn, f"sync_pull_cursor:{owner}")}


@app.post("/sync/cursor")
async def sync_cursor_set(request: Request, service: FinanceService = Depends(get_service)):
    owner = _require_cloud_owner()
    payload = await request.json()
    cursor = str(payload.get("cursor") or "").strip() if isinstance(payload, dict) else ""
    if not cursor or _parse_sync_datetime(cursor) is None:
        raise HTTPException(status_code=422, detail="Cursor de sincronizacion invalido.")
    with service.db.connect() as conn:
        _config_set(conn, f"sync_pull_cursor:{owner}", cursor)
    return {"ok": True, "cursor": cursor}


@app.post("/sync/mark-synced")
async def sync_mark_synced(request: Request, service: FinanceService = Depends(get_service)):
    owner = _require_cloud_owner()
    payload = await request.json()
    accepted = {table: [str(v) for v in payload.get(table, []) if v] for table in SYNC_TABLES}
    now = datetime.now().isoformat(timespec="seconds")

    with service.db.connect() as conn:
        _ensure_sync_metadata_columns(conn)
        for table, sync_ids in accepted.items():
            for sync_id in sync_ids:
                conn.execute(
                    f"""
                    UPDATE {table}
                    SET sync_status = 'synced',
                        last_synced_at = ?
                    WHERE sync_id = ? AND owner_user_id = ?
                    """,
                    (now, sync_id, owner),
                )
    return {"ok": True, "marked": {table: len(values) for table, values in accepted.items()}}


async def _parse_remote_payload(request: Request) -> tuple[dict[str, Any], int, dict[str, int]]:
    if "application/json" not in request.headers.get("content-type", "").lower():
        raise HTTPException(status_code=415, detail="El Content-Type debe ser application/json.")
    raw_body = await request.body()
    try:
        payload = json.loads(raw_body)
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="El body JSON para aplicar datos remotos no es valido.") from exc
    if not isinstance(payload, dict):
        raise HTTPException(status_code=422, detail="El body para aplicar datos remotos debe ser un objeto JSON.")
    counts: dict[str, int] = {}
    for table in SYNC_TABLES:
        items = payload.get(table, [])
        if items is None:
            items = []
        if not isinstance(items, list):
            raise HTTPException(status_code=422, detail=f"El campo {table} debe ser una lista.")
        payload[table] = items
        counts[table] = len(items)
    return payload, len(raw_body), counts


@app.post("/sync/apply-remote/check")
async def sync_apply_remote_check(request: Request):
    owner = _require_cloud_owner()
    _, body_size_bytes, counts = await _parse_remote_payload(request)
    return {
        "ok": True,
        "token_present": bool(request.headers.get(_LOCAL_TOKEN_HEADER)),
        "owner_present": bool(owner),
        "body_size_bytes": body_size_bytes,
        "body_parsed": True,
        "remote_total": sum(counts.values()),
        "counts_by_entity": counts,
        "message": "El POST local protegido y el body JSON son validos.",
    }


async def _sync_apply_remote_impl(request: Request, service: FinanceService):
    owner = _require_cloud_owner()
    payload, _, _ = await _parse_remote_payload(request)
    now = datetime.now().isoformat(timespec="seconds")
    result: dict[str, Any] = {
        f"{table}_{suffix}": 0
        for table in SYNC_TABLES
        for suffix in ("inserted", "updated", "skipped", "kept_local")
    }
    result["conflicts_total"] = 0
    result["applied_remote_total"] = 0
    result["kept_local_total"] = 0
    result["remote_changes_total"] = sum(len(payload.get(table, []) or []) for table in SYNC_TABLES)
    result["conflicts_by_table"] = {}

    with service.db.connect() as conn:
        _ensure_sync_metadata_columns(conn)
        _ensure_sync_history_table(conn)
        current_device = _get_or_create_device_info(conn)
        current_device_id = current_device["device_id"]
        for item in payload.get("categorias", []) or []:
            sync_id = str(item.get("sync_id") or "").strip()
            nombre = str(item.get("nombre") or "").strip()
            tipo = str(item.get("tipo") or "").strip()
            if not sync_id or not nombre or tipo not in {"ingreso", "gasto", "ahorro", "inversion"}:
                continue

            local = conn.execute("SELECT * FROM categorias WHERE sync_id = ? AND owner_user_id = ?", (sync_id, owner)).fetchone()
            if local:
                conflict = _basic_conflict_detected(local, item, current_device_id)
                if conflict and not _remote_is_newer(item, local):
                    _record_sync_conflict(conn, "categorias", local, item, "kept_local", now, result)
                    result["categorias_kept_local"] += 1
                    result["kept_local_total"] += 1
                    continue
                if _remote_is_newer(item, local):
                    if conflict:
                        _record_sync_conflict(conn, "categorias", local, item, "applied_remote", now, result)
                    conn.execute(
                        """
                        UPDATE categorias
                        SET nombre = ?, tipo = ?, updated_at = ?, deleted_at = ?, sync_status = 'synced', last_synced_at = ?,
                            last_remote_device_id = ?, last_remote_device_name = ?, last_remote_updated_at = ?
                        WHERE sync_id = ? AND owner_user_id = ?
                        """,
                        (nombre, tipo, item.get("updated_at") or now, item.get("deleted_at"), now, *_remote_meta_values(item), sync_id, owner),
                    )
                    result["categorias_updated"] += 1
                    result["applied_remote_total"] += 1
                continue

            try:
                conn.execute(
                    """
                    INSERT INTO categorias (
                        nombre, tipo, owner_user_id, sync_id, created_at, updated_at, deleted_at, sync_status, last_synced_at,
                        last_remote_device_id, last_remote_device_name, last_remote_updated_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, 'synced', ?, ?, ?, ?)
                    """,
                    (nombre, tipo, owner, sync_id, item.get("created_at") or now, item.get("updated_at") or now, item.get("deleted_at"), now, *_remote_meta_values(item)),
                )
                result["categorias_inserted"] += 1
                result["applied_remote_total"] += 1
            except sqlite3.IntegrityError:
                existing = conn.execute(
                    "SELECT * FROM categorias WHERE nombre = ? AND tipo = ? AND owner_user_id = ?",
                    (nombre, tipo, owner),
                ).fetchone()
                if existing and _remote_is_newer(item, existing):
                    conn.execute(
                        """
                        UPDATE categorias
                        SET sync_id = ?, updated_at = ?, deleted_at = ?, sync_status = 'synced', last_synced_at = ?,
                            last_remote_device_id = ?, last_remote_device_name = ?, last_remote_updated_at = ?
                        WHERE id = ? AND owner_user_id = ?
                        """,
                        (sync_id, item.get("updated_at") or now, item.get("deleted_at"), now, *_remote_meta_values(item), existing["id"], owner),
                    )
                    result["categorias_updated"] += 1
                    result["applied_remote_total"] += 1

        for item in payload.get("tags", []) or []:
            _apply_remote_tag(conn, item, result, now, current_device_id)

        for item in payload.get("metas_ahorro", []) or []:
            _apply_simple_remote(
                conn,
                "metas_ahorro",
                item,
                ["nombre", "monto_objetivo", "monto_inicial", "fecha_objetivo", "descripcion", "estado"],
                result,
                now,
                current_device_id,
            )

        for table, fields in {
            "gastos_fijos": ["descripcion", "monto", "dia_vencimiento", "activo"],
            "gastos_programados": ["descripcion", "monto_estimado", "fecha_vencimiento", "estado", "es_recurrente", "frecuencia"],
            "presupuestos": ["mes", "anio", "monto"],
        }.items():
            for item in payload.get(table, []) or []:
                categoria_id = _local_category_id(conn, item)
                if categoria_id is None:
                    result[f"{table}_skipped"] = result.get(f"{table}_skipped", 0) + 1
                    continue
                local_item = dict(item)
                local_item["categoria_id"] = categoria_id
                try:
                    _apply_simple_remote(conn, table, local_item, ["categoria_id", *fields], result, now, current_device_id)
                except sqlite3.IntegrityError:
                    result[f"{table}_skipped"] = result.get(f"{table}_skipped", 0) + 1

        for item in payload.get("movimientos", []) or []:
            sync_id = str(item.get("sync_id") or "").strip()
            categoria_id = _local_category_id(conn, item)
            if not sync_id or categoria_id is None:
                result["movimientos_skipped"] += 1
                continue

            local = conn.execute("SELECT * FROM movimientos WHERE sync_id = ? AND owner_user_id = ?", (sync_id, owner)).fetchone()
            values = (
                item.get("fecha"),
                item.get("tipo"),
                categoria_id,
                item.get("descripcion") or "",
                float(item.get("monto") or 0),
                item.get("updated_at") or now,
                item.get("deleted_at"),
                now,
                sync_id,
            )
            if local:
                conflict = _basic_conflict_detected(local, item, current_device_id)
                if conflict and not _remote_is_newer(item, local):
                    _record_sync_conflict(conn, "movimientos", local, item, "kept_local", now, result)
                    result["movimientos_kept_local"] += 1
                    result["kept_local_total"] += 1
                    continue
                if _remote_is_newer(item, local):
                    if conflict:
                        _record_sync_conflict(conn, "movimientos", local, item, "applied_remote", now, result)
                    conn.execute(
                        """
                        UPDATE movimientos
                        SET fecha = ?, tipo = ?, categoria_id = ?, descripcion = ?, monto = ?,
                            updated_at = ?, deleted_at = ?, sync_status = 'synced', last_synced_at = ?,
                            last_remote_device_id = ?, last_remote_device_name = ?, last_remote_updated_at = ?
                        WHERE sync_id = ? AND owner_user_id = ?
                        """,
                        (*values[:-1], *_remote_meta_values(item), values[-1], owner),
                    )
                    result["movimientos_updated"] += 1
                    result["applied_remote_total"] += 1
                continue

            conn.execute(
                """
                INSERT INTO movimientos (
                    fecha, tipo, categoria_id, descripcion, monto,
                    sync_id, owner_user_id, created_at, updated_at, deleted_at, sync_status, last_synced_at,
                    last_remote_device_id, last_remote_device_name, last_remote_updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'synced', ?, ?, ?, ?)
                """,
                (
                    item.get("fecha"),
                    item.get("tipo"),
                    categoria_id,
                    item.get("descripcion") or "",
                    float(item.get("monto") or 0),
                    sync_id,
                    owner,
                    item.get("created_at") or now,
                    item.get("updated_at") or now,
                    item.get("deleted_at"),
                    now,
                    *_remote_meta_values(item),
                ),
            )
            result["movimientos_inserted"] += 1
            result["applied_remote_total"] += 1

        for item in payload.get("movimiento_tags", []) or []:
            _apply_remote_movement_tag(conn, item, result, now)

    applied = {table: int(result.get(f"{table}_inserted", 0)) + int(result.get(f"{table}_updated", 0)) for table in SYNC_TABLES}
    kept_local = {table: int(result.get(f"{table}_kept_local", 0)) for table in SYNC_TABLES}
    return {
        "ok": True,
        "result": result,
        "applied": applied,
        "kept_local": kept_local,
        "conflicts": {
            "total": int(result.get("conflicts_total", 0)),
            "by_table": result.get("conflicts_by_table", {}),
        },
        "remote_changes_total": int(result.get("remote_changes_total", 0)),
        "applied_remote_total": int(result.get("applied_remote_total", 0)),
        "kept_local_total": int(result.get("kept_local_total", 0)),
    }


@app.post("/sync/apply-remote")
async def sync_apply_remote(request: Request, service: FinanceService = Depends(get_service)):
    try:
        return await _sync_apply_remote_impl(request, service)
    except HTTPException:
        raise
    except Exception as exc:
        _logger.exception("Error aplicando datos remotos. owner=%s error_type=%s", _safe_owner_for_log(_current_owner()), type(exc).__name__)
        raise HTTPException(status_code=500, detail="No se pudieron aplicar los datos remotos en el servicio local.") from exc


@app.post("/backup/export")
def export_backup(service: FinanceService = Depends(get_service)):
    source = Path(service.db.db_path)
    if not source.exists():
        raise HTTPException(status_code=404, detail="No existe la base de datos.")
    # SQLite Backup API genera una copia consistente incluso si hay escrituras
    # concurrentes; copy2 sobre una DB viva podia producir exports incompletos.
    output = service.backup_database(Path(mkdtemp(prefix="finanzas_backup_")))
    # Cleanup best-effort tras streaming; startup recolecta temporales antiguos
    # si una descarga interrumpida impide ejecutar BackgroundTask.
    return FileResponse(
        path=output,
        filename=output.name,
        media_type="application/octet-stream",
        background=BackgroundTask(_cleanup_temp_snapshot, output),
    )


@app.post("/backup/restore")
async def restore_backup(request: Request, service: FinanceService = Depends(get_service)):
    try:
        ensure_app_data_initialized()
        try:
            payload = await request.json()
        except Exception as exc:
            _logger.exception("Body invalido en restauracion de copia. error_type=%s", type(exc).__name__)
            raise HTTPException(status_code=400, detail="Solicitud invalida para restaurar copia de seguridad.") from exc

        _logger.info("Payload recibido en /backup/restore: keys=%s", list(payload.keys()) if isinstance(payload, dict) else type(payload).__name__)
        if not isinstance(payload, dict):
            raise HTTPException(status_code=400, detail="Solicitud invalida para restaurar copia de seguridad.")

        source_path_value = str(payload.get("source_path") or "").strip()
        if not source_path_value:
            raise HTTPException(status_code=400, detail="Debes indicar source_path.")

        source = Path(source_path_value).expanduser()
        pre_restore_backups = get_data_dir() / "backups"
        safety = service.restore_database_from_path(source, pre_restore_backups)
        invalidate_app_data_initialized()
        _logger.info("Restaurar copia OK. source=%s backup_pre_restore=%s", source.name, safety.name)
        return {"ok": True, "safety_backup": str(safety)}
    except HTTPException:
        raise
    except (FileNotFoundError, ValueError) as exc:
        # No registrar paths completos provenientes del request de restore.
        _logger.exception("Error restaurando copia de seguridad. error_type=%s", type(exc).__name__)
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:
        _logger.exception("Error restaurando copia de seguridad. error_type=%s", type(exc).__name__)
        raise HTTPException(
            status_code=500,
            detail="No se pudo restaurar la copia de seguridad.",
        ) from exc


@app.get("/backup/download")
def download_backup(service: FinanceService = Depends(get_service)):
    try:
        ensure_app_data_initialized()
        source = Path(service.db.db_path)
        _logger.info("Solicitud de copia de seguridad. db_file=%s", source.name)
        if not source.exists():
            _logger.error("No existe la DB para copia de seguridad. db_file=%s", source.name)
            raise HTTPException(status_code=404, detail="No existe la base de datos local.")
        if not source.is_file():
            _logger.error("La ruta de DB no es un archivo. db_file=%s", source.name)
            raise HTTPException(status_code=500, detail="No se pudo obtener la copia de seguridad.")
        if source.stat().st_size <= 0:
            _logger.error("La DB esta vacia. db_file=%s", source.name)
            raise HTTPException(status_code=500, detail="No se pudo obtener la copia de seguridad.")

        filename = f"ScisoNomics_copia_seguridad_{datetime.now().strftime('%Y-%m-%d')}.db"
        # Servir un snapshot consistente: entregar la DB viva podia omitir
        # paginas en escritura mientras el usuario descargaba la copia.
        snapshot = service.backup_database(Path(mkdtemp(prefix="finanzas_download_")))
        # Cleanup best-effort tras streaming; startup recolecta temporales antiguos
        # si una descarga interrumpida impide ejecutar BackgroundTask.
        return FileResponse(
            path=snapshot,
            filename=filename,
            media_type="application/octet-stream",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
            background=BackgroundTask(_cleanup_temp_snapshot, snapshot),
        )
    except HTTPException:
        raise
    except Exception as exc:
        _logger.exception("Error generando copia de seguridad. error_type=%s", type(exc).__name__)
        raise HTTPException(status_code=500, detail="No se pudo obtener la copia de seguridad.") from exc


@app.get("/export/excel")
def export_excel(
    month: int | None = Query(default=None, ge=1, le=12),
    year: int | None = Query(default=None),
    desde: str | None = Query(default=None),
    hasta: str | None = Query(default=None),
    service: FinanceService = Depends(get_service),
):
    now = datetime.now()
    resolved_month = month if month is not None else now.month
    resolved_year = year if year is not None else now.year
    output = Path(gettempdir()) / "ScisoNomics" / "exports" / f"ScisoNomics_{now.strftime('%Y-%m-%d')}.xlsx"
    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        if desde or hasta:
            if not desde or not hasta:
                raise HTTPException(status_code=400, detail="Debes indicar ambas fechas: desde y hasta.")
            try:
                from_dt = datetime.strptime(desde, "%Y-%m-%d")
                to_dt = datetime.strptime(hasta, "%Y-%m-%d")
            except ValueError as exc:
                raise HTTPException(status_code=400, detail="Formato de fecha inválido. Usa YYYY-MM-DD.") from exc
            if from_dt > to_dt:
                raise HTTPException(status_code=400, detail="La fecha 'desde' no puede ser mayor que 'hasta'.")
            export_date_range_report(service, desde, hasta, output)
        else:
            export_monthly_report(service, resolved_month, resolved_year, output)
    except HTTPException:
        raise
    except Exception as exc:
        _logger.exception("Error exportando Excel. month=%s year=%s error_type=%s", resolved_month, resolved_year, type(exc).__name__)
        raise HTTPException(status_code=500, detail="No se pudo generar el archivo Excel.") from exc
    return FileResponse(
        path=output,
        filename=output.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{output.name}"'},
    )


@app.get("/export/mensual")
def export_mensual(month: int = Query(..., ge=1, le=12), year: int = Query(...), service: FinanceService = Depends(get_service)):
    output = Path(mkdtemp(prefix="finanzas_export_")) / f"movimientos_{year}_{month:02d}.xlsx"
    export_monthly_report(service, month, year, output)
    return FileResponse(path=output, filename=output.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/export/anual")
def export_anual(year: int = Query(...), service: FinanceService = Depends(get_service)):
    output = Path(mkdtemp(prefix="finanzas_export_")) / f"finanzas_{year}.xlsx"
    export_yearly_report(service, year, output)
    return FileResponse(path=output, filename=output.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/export/filtrado")
def export_filtrado(
    month: int = Query(..., ge=1, le=12),
    year: int = Query(...),
    tipo: Literal["todos", "ingreso", "gasto", "ahorro", "inversion"] = "todos",
    search: str = "",
    service: FinanceService = Depends(get_service),
):
    current_type = None if tipo in ("todos", "inversion") else tipo
    rows = service.list_movimientos(month, year, current_type)
    if tipo == "inversion":
        rows = [r for r in rows if "invers" in str(r.get("categoria", "")).lower()]
    term = search.strip().lower()
    if term:
        rows = [
            row
            for row in rows
            if term in str(row.get("descripcion", "")).lower()
            or term in str(row.get("categoria", "")).lower()
            or term in str(row.get("fecha", "")).lower()
        ]
    output = Path(mkdtemp(prefix="finanzas_export_")) / f"filtrado_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    export_filtered_movimientos(rows, output)
    return FileResponse(path=output, filename=output.name, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _backup_dir() -> Path:
    return WEB_DB_PATH.parent.parent / "backups"


@app.get("/backups")
def list_backups(service: FinanceService = Depends(get_service)):
    items = service.list_backups(_backup_dir())
    return {
        "folder": str(_backup_dir()),
        "count": len(items),
        "last_backup": items[0] if items else None,
        "items": items,
        "frequency": service.get_config_value("backup_frequency", "desactivado"),
    }


@app.post("/backups/create")
def create_backup(service: FinanceService = Depends(get_service)):
    created = service.backup_database(_backup_dir())
    return {"ok": True, "file": str(created), "name": created.name}


@app.post("/backups/restore")
def restore_backup_from_name(payload: BackupRestoreIn, service: FinanceService = Depends(get_service)):
    if Path(payload.file_name).name != payload.file_name:
        raise HTTPException(status_code=400, detail="Nombre de backup invalido.")
    backup_file = _backup_dir() / payload.file_name
    safety = service.restore_database(backup_file, _backup_dir())
    invalidate_app_data_initialized()
    return {"ok": True, "safety_backup": str(safety)}


@app.post("/backups/frequency")
def set_backup_frequency(payload: BackupFrequencyIn, service: FinanceService = Depends(get_service)):
    service.set_config_value("backup_frequency", payload.frecuencia)
    return {"ok": True}


@app.get("/metas")
def list_metas(service: FinanceService = Depends(get_service)):
    return service.list_metas_ahorro()


@app.post("/metas")
def create_meta(payload: MetaAhorroIn, service: FinanceService = Depends(get_service)):
    try:
        service.create_meta_ahorro(MetaAhorroInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "La meta") from exc


@app.put("/metas/{meta_id}")
def update_meta(meta_id: int, payload: MetaAhorroIn, service: FinanceService = Depends(get_service)):
    try:
        service.update_meta_ahorro(meta_id, MetaAhorroInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "La meta") from exc


@app.delete("/metas/{meta_id}")
def delete_meta(meta_id: int, service: FinanceService = Depends(get_service)):
    try:
        service.delete_meta_ahorro(meta_id)
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "La meta") from exc


@app.get("/tags")
def list_tags(service: FinanceService = Depends(get_service)):
    return service.list_tags()


@app.post("/tags")
def create_tag(payload: TagIn, service: FinanceService = Depends(get_service)):
    try:
        service.create_tag(TagInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "La etiqueta") from exc


@app.put("/tags/{tag_id}")
def update_tag(tag_id: int, payload: TagIn, service: FinanceService = Depends(get_service)):
    try:
        service.update_tag(tag_id, TagInput(**payload.model_dump()))
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "La etiqueta") from exc


@app.delete("/tags/{tag_id}")
def delete_tag(tag_id: int, service: FinanceService = Depends(get_service)):
    try:
        service.delete_tag(tag_id)
        return {"ok": True}
    except ValueError as exc:
        raise _service_value_error(exc, "La etiqueta") from exc


@app.get("/calendario")
def calendario(
    month: int | None = Query(default=None, ge=1, le=12),
    year: int | None = Query(default=None),
    mes: int | None = Query(default=None, ge=1, le=12),
    anio: int | None = Query(default=None),
    service: FinanceService = Depends(get_service),
):
    resolved_month = month if month is not None else mes
    resolved_year = year if year is not None else anio
    if resolved_month is None or resolved_year is None:
        raise HTTPException(status_code=400, detail="Debes indicar month/year o mes/anio.")
    return service.get_calendario_mensual(resolved_month, resolved_year)


@app.get("/reporte-mensual")
def reporte_mensual(
    month: int | None = Query(default=None, ge=1, le=12),
    year: int | None = Query(default=None),
    mes: int | None = Query(default=None, ge=1, le=12),
    anio: int | None = Query(default=None),
    service: FinanceService = Depends(get_service),
):
    resolved_month = month if month is not None else mes
    resolved_year = year if year is not None else anio
    if resolved_month is None or resolved_year is None:
        raise HTTPException(status_code=400, detail="Debes indicar month/year o mes/anio.")

    raw = service.get_reporte_mensual_avanzado(resolved_month, resolved_year) or {}
    return {
        "month": resolved_month,
        "year": resolved_year,
        "ingresos": float(raw.get("ingresos_mes", 0) or 0),
        "gastos": float(raw.get("gastos_mes", 0) or 0),
        "ahorro": float(raw.get("ahorro_mes", 0) or 0),
        "inversiones": float(raw.get("inversiones_mes", 0) or 0),
        "balance_operativo": float(raw.get("balance_operativo", 0) or 0),
        "disponible_luego_ahorro": float(raw.get("disponible_luego_ahorro", 0) or 0),
        "top_categorias": list(raw.get("top_5_categorias_gasto", []) or []),
        "top_movimientos": list(raw.get("top_5_movimientos_gasto", []) or []),
        "evolucion_ultimos_6_meses": list(raw.get("evolucion_6_meses", []) or []),
        "presupuestos_excedidos": list(raw.get("presupuestos_excedidos", []) or []),
        "metas": list(raw.get("metas_activas", []) or []),
    }


@app.post("/import/preview")
async def import_preview(file: UploadFile = File(...)):
    name = (file.filename or "").lower()
    content = await file.read()
    if name.endswith(".csv"):
        text = content.decode("utf-8", errors="replace")
        rows = list(csv.DictReader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        values = list(ws.values)
        if not values:
            rows = []
        else:
            headers = [str(h or "").strip() for h in values[0]]
            rows = [{headers[i]: r[i] for i in range(len(headers))} for r in values[1:]]
    else:
        raise HTTPException(status_code=400, detail="Formato no soportado. Usa .csv o .xlsx")
    return {"rows_read": len(rows), "preview": rows[:50]}

